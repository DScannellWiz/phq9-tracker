import argparse
import os
import sqlite3
import subprocess
import sys
import tempfile
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from xml.sax.saxutils import escape
from tkinter import (
    BOTH,
    BOTTOM,
    END,
    LEFT,
    RIGHT,
    TOP,
    Button,
    Canvas,
    Checkbutton,
    Entry,
    Frame,
    IntVar,
    Label,
    LabelFrame,
    Menu,
    StringVar,
    Text,
    Tk,
    Toplevel,
    filedialog,
    messagebox,
    ttk,
)

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
except Exception:
    colors = None
    inch = 72
    Image = None
    PageBreak = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None
    Table = None
    TableStyle = None

try:
    from PIL import Image as PILImage
    from PIL import ImageDraw, ImageFont
except Exception:
    PILImage = None


APP_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = APP_DIR.parents[1] if APP_DIR.parent.name == "src" else APP_DIR.parent
DATA_DIR = PROJECT_ROOT / "data"
EXPORTS_DIR = PROJECT_ROOT / "exports"
REPORTS_DIR = PROJECT_ROOT / "reports"
ICON_PATH = PROJECT_ROOT / "packaging" / "assets" / "PHQ9_Tracker.ico"


def default_db_path() -> Path:
    explicit = os.environ.get("PHQ9_TRACKER_DB_PATH")
    if explicit:
        return Path(explicit)
    if getattr(sys, "frozen", False):
        if os.environ.get("PHQ9_TRACKER_PORTABLE") == "1":
            return Path(sys.executable).resolve().parent / "phq9_tracker.sqlite"
        local_app_data = os.environ.get("LOCALAPPDATA")
        if local_app_data:
            return Path(local_app_data) / "PHQ9Tracker" / "phq9_tracker.sqlite"
    return DATA_DIR / "phq9_tracker.sqlite"


DB_PATH = default_db_path()
BUNDLED_PYTHON = Path(os.environ["PHQ9_TRACKER_BUNDLED_PYTHON"]) if os.environ.get("PHQ9_TRACKER_BUNDLED_PYTHON") else None
DISCLAIMER = "This report is for discussion with a licensed clinician and is not a diagnosis."
DAILY_SCORE_LABEL = "Daily Severity Score"
FREQUENCY_SCORE_LABEL = "14-Day Symptom Frequency Score"
ANALYSIS_WORKBOOK_SCHEMA_VERSION = "1.0"
SCORING_EXPLANATION = """Mental Health Tracker calculates two related but different measurements.

Daily Severity Score
Each PHQ-9 or GAD-7 item is rated from 0 (symptom not present) to 3 (high symptom severity). The item responses are summed for that date. This answers: How severe were the reported symptoms on this particular day?

14-Day Symptom Frequency Score
For the 14 calendar days ending on the selected date, each response greater than 0 counts as one symptom-present day. For each item, 0 days = 0 points, 1-6 days = 1 point, 7-11 days = 2 points, and 12-14 days = 3 points. The converted item scores are summed. This answers: How consistently were these symptoms present during the last 14 days?

Important distinction
A daily response of 1 and a daily response of 3 each count as one symptom-present day in the 14-day calculation, although they contribute differently to the Daily Severity Score. For example, an item present on 6 days receives 1 frequency point; an item present on 12 days receives 3 frequency points.

Data coverage
The application displays how many daily entries were available. Calendar days without an entry are treated as days with no recorded symptom-present response; that is not proof that the symptom was absent.

Mindful check-ins
Daily check-ins should encourage mindful reflection rather than rapid completion. The application intentionally requires users to consider each symptom individually to reduce habitual responses and improve the quality of the recorded data."""

TREATMENT_EVENT_TYPES = [
    "Therapy",
    "Ketamine infusion",
    "Psychiatry appointment",
    "Medication change",
    "Primary care appointment",
    "Exercise",
    "Support group",
]

PHQ9_ITEM_LABELS = [
    "Little interest or pleasure in doing things",
    "Feeling down, depressed, or hopeless",
    "Sleep changes",
    "Low energy",
    "Appetite changes",
    "Feeling bad about yourself",
    "Trouble concentrating",
    "Moving/speaking slowly or being restless",
    "Thoughts of self-harm",
]

GAD7_ITEM_LABELS = [
    "Feeling nervous, anxious, or on edge",
    "Not being able to stop or control worrying",
    "Worrying too much about different things",
    "Trouble relaxing",
    "Being so restless that it is hard to sit still",
    "Becoming easily annoyed or irritable",
    "Feeling afraid as if something awful might happen",
]

ITEM_LABELS = PHQ9_ITEM_LABELS


@dataclass(frozen=True)
class AssessmentDefinition:
    assessment_id: str
    display_name: str
    short_name: str
    item_labels: list[str]
    max_score: int

    @property
    def item_count(self) -> int:
        return len(self.item_labels)


ASSESSMENTS = {
    "phq9": AssessmentDefinition("phq9", "PHQ-9", "PHQ-9", PHQ9_ITEM_LABELS, 27),
    "gad7": AssessmentDefinition("gad7", "GAD-7", "GAD-7", GAD7_ITEM_LABELS, 21),
}
ASSESSMENT_ORDER = ["phq9", "gad7"]

@dataclass
class EntryRow:
    id: int
    entry_date: str
    items: list[int]
    total: int
    severity: str
    notes: str
    note_tag: str = ""


@dataclass
class AssessmentEntryRow:
    id: int
    assessment_id: str
    entry_date: str
    items: list[int]
    total: int
    severity: str
    notes: str
    note_tag: str = ""


@dataclass
class FourteenDayScore:
    item_counts: list[int]
    item_scores: list[int]
    total_score: int
    severity: str
    start_date: str | None
    end_date: str | None
    entries_included: int
    calendar_days: int = 14


@dataclass(frozen=True)
class PeriodComparison:
    assessment_id: str
    current: FourteenDayScore
    previous: FourteenDayScore
    current_start: str
    current_end: str
    previous_start: str
    previous_end: str

    @property
    def delta(self) -> int:
        return self.current.total_score - self.previous.total_score

    @property
    def has_comparable_data(self) -> bool:
        return self.current.entries_included > 0 and self.previous.entries_included > 0


@dataclass(frozen=True)
class TreatmentCycle:
    """A descriptive window anchored to a recorded ketamine infusion."""

    label: str
    start_date: str
    end_date: str
    entries: list[EntryRow | AssessmentEntryRow]


@dataclass(frozen=True)
class ChartPoint:
    """One discoverable score point rendered on a Review chart."""

    x: float
    y: float
    entry_date: str
    series_name: str
    score: int
    color: str


@dataclass(frozen=True)
class Item9Context:
    """Recency and coverage facts for PHQ-9 item 9 discussion language."""

    recent_start: str
    recent_end: str
    recent_checkins: int
    recent_above_zero: int
    historical_above_zero: int
    latest_historical_date: str | None

    @property
    def has_recent_above_zero(self) -> bool:
        return self.recent_above_zero > 0

    @property
    def has_historical_above_zero(self) -> bool:
        return self.historical_above_zero > 0


def severity_for_score(score: int) -> str:
    if score <= 4:
        return "Minimal"
    if score <= 9:
        return "Mild"
    if score <= 14:
        return "Moderate"
    if score <= 19:
        return "Moderately severe"
    return "Severe"


def gad7_severity_for_score(score: int) -> str:
    if score <= 4:
        return "Minimal"
    if score <= 9:
        return "Mild"
    if score <= 14:
        return "Moderate"
    return "Severe"


def severity_for_assessment(assessment_id: str, score: int) -> str:
    if assessment_id == "gad7":
        return gad7_severity_for_score(score)
    return severity_for_score(score)


def convert_14_day_count_to_item_score(days_present: int) -> int:
    if days_present <= 0:
        return 0
    if days_present <= 6:
        return 1
    if days_present <= 11:
        return 2
    return 3


def calculate_14_day_symptom_frequency_score(
    entries: list[EntryRow | AssessmentEntryRow],
    item_count: int | None = None,
    assessment_id: str = "phq9",
) -> FourteenDayScore:
    """Calculate the clinical-style 14-day symptom-frequency score.

    The window is the most recent 14 calendar days ending on the most recent
    entry date. For each item, each day with an entry and item score > 0 counts
    as one symptom-present day. Missing calendar days are treated as no recorded
    symptom-present day for the count, and entries_included reports how many
    actual entry rows were available in the window.
    """
    item_count = item_count or (len(entries[0].items) if entries else ASSESSMENTS[assessment_id].item_count)
    if not entries:
        return FourteenDayScore(
            [0] * item_count,
            [0] * item_count,
            0,
            severity_for_assessment(assessment_id, 0),
            None,
            None,
            0,
        )
    sorted_entries = sorted(entries, key=lambda row: row.entry_date)
    end_dt = datetime.fromisoformat(sorted_entries[-1].entry_date).date()
    start_dt = end_dt - timedelta(days=13)
    window_entries = [
        row
        for row in sorted_entries
        if start_dt <= datetime.fromisoformat(row.entry_date).date() <= end_dt
    ]
    item_counts = []
    for idx in range(item_count):
        item_counts.append(sum(1 for row in window_entries if idx < len(row.items) and row.items[idx] > 0))
    item_scores = [convert_14_day_count_to_item_score(count) for count in item_counts]
    total_score = sum(item_scores)
    return FourteenDayScore(
        item_counts=item_counts,
        item_scores=item_scores,
        total_score=total_score,
        severity=severity_for_assessment(assessment_id, total_score),
        start_date=start_dt.isoformat(),
        end_date=end_dt.isoformat(),
        entries_included=len(window_entries),
    )


def calculate_symptom_frequency_score_for_window(
    entries: list[EntryRow | AssessmentEntryRow],
    window_start: str,
    window_end: str,
    item_count: int,
    assessment_id: str,
) -> FourteenDayScore:
    """Calculate a symptom-frequency score for an explicit calendar window."""
    start_dt = datetime.fromisoformat(window_start).date()
    end_dt = datetime.fromisoformat(window_end).date()
    if end_dt < start_dt:
        raise ValueError("Window end date must not be before its start date.")
    window_entries = [
        row
        for row in entries
        if start_dt <= datetime.fromisoformat(row.entry_date).date() <= end_dt
    ]
    item_counts = [
        sum(1 for row in window_entries if idx < len(row.items) and row.items[idx] > 0)
        for idx in range(item_count)
    ]
    item_scores = [convert_14_day_count_to_item_score(count) for count in item_counts]
    total_score = sum(item_scores)
    return FourteenDayScore(
        item_counts=item_counts,
        item_scores=item_scores,
        total_score=total_score,
        severity=severity_for_assessment(assessment_id, total_score),
        start_date=window_start,
        end_date=window_end,
        entries_included=len(window_entries),
        calendar_days=(end_dt - start_dt).days + 1,
    )


def compare_recent_14_day_periods(
    entries: list[EntryRow | AssessmentEntryRow],
    assessment_id: str,
    end_date: str,
) -> PeriodComparison:
    """Compare the latest 14 calendar days with the immediately prior 14 days."""
    definition = ASSESSMENTS[assessment_id]
    current_end = datetime.fromisoformat(end_date).date()
    current_start = current_end - timedelta(days=13)
    previous_end = current_start - timedelta(days=1)
    previous_start = previous_end - timedelta(days=13)
    current = calculate_symptom_frequency_score_for_window(
        entries, current_start.isoformat(), current_end.isoformat(), definition.item_count, assessment_id
    )
    previous = calculate_symptom_frequency_score_for_window(
        entries, previous_start.isoformat(), previous_end.isoformat(), definition.item_count, assessment_id
    )
    return PeriodComparison(
        assessment_id=assessment_id,
        current=current,
        previous=previous,
        current_start=current_start.isoformat(),
        current_end=current_end.isoformat(),
        previous_start=previous_start.isoformat(),
        previous_end=previous_end.isoformat(),
    )


def comparison_trend_text(comparison: PeriodComparison, include_arrow: bool = False) -> str:
    """Return neutral, non-diagnostic wording for a two-period comparison."""
    if not comparison.has_comparable_data:
        return "Not enough data for comparison"
    if comparison.delta < 0:
        return f"{'Down - ' if include_arrow else ''}Lower by {abs(comparison.delta)} points"
    if comparison.delta > 0:
        return f"{'Up - ' if include_arrow else ''}Higher by {comparison.delta} points"
    return f"{'Right - ' if include_arrow else ''}No score change"


def shift_calendar_date(value: str, days: int, today: date | None = None) -> str:
    """Move an ISO date by whole calendar days without allowing future dates."""
    parsed = parse_date(value)
    if not parsed:
        raise ValueError("Enter the date as YYYY-MM-DD.")
    shifted = datetime.fromisoformat(parsed).date() + timedelta(days=days)
    if shifted > (today or date.today()):
        raise ValueError("Future check-ins are not available.")
    return shifted.isoformat()


def validate_nonfuture_date(value: str, today: date | None = None) -> str:
    """Normalize an ISO date and reject dates after the local current day."""
    parsed = parse_date(value)
    if not parsed:
        raise ValueError("Enter the date as YYYY-MM-DD.")
    if datetime.fromisoformat(parsed).date() > (today or date.today()):
        raise ValueError("Future dates are not available.")
    return parsed


def history_record_count(day_data: dict) -> int:
    """Count the independently manageable records represented by one day."""
    return len(day_data["assessments"]) + len(day_data["events"]) + int(bool(day_data["notes"] or day_data["note_tag"]))


def history_status_text(entry_date: str, day_data: dict) -> str:
    """Return calm History status text without implying that a record failed to load."""
    record_count = history_record_count(day_data)
    if record_count == 0:
        return (
            f"No existing records for {entry_date}. History can edit only dates that already contain "
            "a check-in, note, or treatment event."
        )
    noun = "record" if record_count == 1 else "records"
    return f"Showing {record_count} existing {noun} for {entry_date}. Edits apply only to this loaded date."


def responsive_canvas_dimension(actual: int, configured: int) -> int:
    """Use live canvas dimensions after layout, including when the window shrinks."""
    return actual if actual > 1 else configured


def chart_vertical_positions(title_bottom: int) -> tuple[int, int]:
    """Keep the legend below the rendered title and the plot below the legend."""
    legend_top = title_bottom + 8
    plot_top = legend_top + 24
    return legend_top, plot_top


def chart_point_callout_text(point: ChartPoint) -> str:
    """Return the exact date and score shown by chart point callouts."""
    return f"{point.entry_date}  |  {point.series_name}: {point.score}"


def nearest_chart_point(
    points: list[ChartPoint],
    x: float,
    y: float,
    max_distance: float = 12,
) -> ChartPoint | None:
    """Find a chart point only when the pointer is within a forgiving hit target."""
    if not points:
        return None
    nearest = min(points, key=lambda point: (point.x - x) ** 2 + (point.y - y) ** 2)
    distance_squared = (nearest.x - x) ** 2 + (nearest.y - y) ** 2
    return nearest if distance_squared <= max_distance**2 else None


def entries_for_window(entries, start_date: str, end_date: str):
    return [row for row in entries if start_date <= row.entry_date <= end_date]


def item9_context(entries: list[AssessmentEntryRow], end_date: str) -> Item9Context:
    """Separate recent item 9 responses from older selected-history context."""
    recent_end = datetime.fromisoformat(end_date).date()
    recent_start = recent_end - timedelta(days=13)
    recent_start_text = recent_start.isoformat()
    recent_end_text = recent_end.isoformat()
    recent_entries = entries_for_window(entries, recent_start_text, recent_end_text)
    historical_positive_entries = [
        row for row in entries if row.entry_date < recent_start_text and row.items[8] > 0
    ]
    return Item9Context(
        recent_start=recent_start_text,
        recent_end=recent_end_text,
        recent_checkins=len(recent_entries),
        recent_above_zero=sum(row.items[8] > 0 for row in recent_entries),
        historical_above_zero=len(historical_positive_entries),
        latest_historical_date=(
            max(row.entry_date for row in historical_positive_entries)
            if historical_positive_entries
            else None
        ),
    )


def item9_context_summary(context: Item9Context) -> str | None:
    """Describe item 9 recency without turning historical data into a current-risk claim."""
    coverage = f"Recent check-in coverage was {context.recent_checkins} of 14 calendar days"
    missing = "days without a check-in are missing information"
    if context.has_recent_above_zero:
        noun = "check-in" if context.recent_above_zero == 1 else "check-ins"
        return (
            f"An above-zero PHQ-9 item 9 response was recorded on {context.recent_above_zero} {noun} "
            f"in the most recent 14-day window ({context.recent_start} to {context.recent_end}). "
            f"{coverage}; {missing}. This is a factual discussion point, not an assessment of current safety."
        )
    if context.has_historical_above_zero:
        noun = "check-in" if context.historical_above_zero == 1 else "check-ins"
        return (
            f"An above-zero PHQ-9 item 9 response was recorded earlier in the selected history on "
            f"{context.historical_above_zero} {noun}, most recently on {context.latest_historical_date}. "
            f"No above-zero item 9 response was recorded among {context.recent_checkins} PHQ-9 check-ins in the "
            f"most recent 14-day window ({context.recent_start} to {context.recent_end}). {coverage}; {missing}. "
            "This older information is neutral historical context and does not establish current risk."
        )
    return None


def item9_conversation_starter(context: Item9Context) -> str | None:
    """Return a recency-aware prompt only when item 9 was above zero in selected history."""
    if context.has_recent_above_zero:
        return (
            f"PHQ-9 item 9 was above zero on {context.recent_above_zero} of {context.recent_checkins} recorded "
            f"check-ins in the most recent 14-day window. Recent check-in coverage was "
            f"{context.recent_checkins} of 14 calendar days; missing days are missing information. "
            "Would it be useful to discuss when this was recorded and what support, if any, would be useful now?"
        )
    if context.has_historical_above_zero:
        return (
            f"An above-zero PHQ-9 item 9 response was recorded earlier in the selected history, most recently on "
            f"{context.latest_historical_date}. It was not recorded above zero among the "
            f"{context.recent_checkins} check-ins in the most recent 14-day window; recent check-in coverage was "
            f"{context.recent_checkins} of 14 calendar days. This is historical context and does not indicate "
            "current risk; would it be useful to discuss what has changed since then?"
        )
    return None


def overall_pattern_summary(
    phq_entries: list[AssessmentEntryRow],
    gad_entries: list[AssessmentEntryRow],
    end_date: str,
) -> str:
    """Describe adjacent 14-day patterns without diagnosis or causal language."""
    phrases = []
    coverage = []
    for assessment_id, entries in (("phq9", phq_entries), ("gad7", gad_entries)):
        comparison = compare_recent_14_day_periods(entries, assessment_id, end_date)
        label = ASSESSMENTS[assessment_id].display_name
        coverage.append(f"{label} {comparison.current.entries_included}/14")
        if not comparison.has_comparable_data:
            phrases.append(f"{label} does not yet have recorded check-ins in both comparison periods")
        elif comparison.delta < 0:
            phrases.append(f"{label} symptom-frequency scores were lower than in the preceding 14 days")
        elif comparison.delta > 0:
            phrases.append(f"{label} symptom-frequency scores were higher than in the preceding 14 days")
        else:
            phrases.append(f"{label} symptom-frequency scores were unchanged from the preceding 14 days")
    return f"{' '.join(f'{phrase}.' for phrase in phrases)} Current-period coverage: {', '.join(coverage)} recorded check-ins."


def symptom_highlights(
    assessment_id: str,
    entries: list[AssessmentEntryRow],
    end_date: str,
    limit: int = 2,
) -> list[str]:
    """Return the largest recorded item-frequency changes using neutral wording."""
    definition = ASSESSMENTS[assessment_id]
    comparison = compare_recent_14_day_periods(entries, assessment_id, end_date)
    current_entries = entries_for_window(entries, comparison.current_start, comparison.current_end)
    previous_entries = entries_for_window(entries, comparison.previous_start, comparison.previous_end)
    if not current_entries:
        return [f"No {definition.display_name} check-ins were recorded in the current 14-day period."]

    current_counts = [sum(row.items[idx] > 0 for row in current_entries) for idx in range(definition.item_count)]
    if not previous_entries:
        ranked = sorted(range(definition.item_count), key=lambda idx: (-current_counts[idx], idx))
        return [
            f"Responses related to {definition.item_labels[idx].lower()} were recorded on {current_counts[idx]} of {len(current_entries)} {definition.display_name} check-ins in this period."
            for idx in ranked[:limit]
            if current_counts[idx] > 0
        ] or [f"No {definition.display_name} symptoms were recorded as present in the current period."]

    previous_counts = [sum(row.items[idx] > 0 for row in previous_entries) for idx in range(definition.item_count)]
    ranked = sorted(
        range(definition.item_count),
        key=lambda idx: (-abs(current_counts[idx] - previous_counts[idx]), idx),
    )
    highlights = []
    for idx in ranked:
        delta = current_counts[idx] - previous_counts[idx]
        if delta == 0:
            continue
        direction = "more often" if delta > 0 else "less often"
        highlights.append(
            f"Responses related to {definition.item_labels[idx].lower()} were recorded {direction}: {current_counts[idx]} of {len(current_entries)} check-ins, compared with {previous_counts[idx]} of {len(previous_entries)} previously."
        )
        if len(highlights) == limit:
            break
    return highlights or [f"Recorded {definition.display_name} symptom frequencies were similar across the two periods."]


def treatment_cycles(
    entries: list[EntryRow | AssessmentEntryRow],
    events: list[tuple[int, str, str, str]],
    end_date: str,
) -> list[TreatmentCycle]:
    """Return current and previous windows between recorded ketamine infusions."""
    anchors = sorted(
        {event_date for _, event_date, event_type, _ in events if normalize_event_type(event_type) == "Ketamine" and event_date <= end_date}
    )
    if not anchors:
        return []
    cycles = []
    latest = anchors[-1]
    cycles.append(TreatmentCycle("Current cycle", latest, end_date, entries_for_window(entries, latest, end_date)))
    if len(anchors) > 1:
        previous = anchors[-2]
        previous_end = (datetime.fromisoformat(latest).date() - timedelta(days=1)).isoformat()
        cycles.append(TreatmentCycle("Previous cycle", previous, previous_end, entries_for_window(entries, previous, previous_end)))
    return cycles


def treatment_cycle_observation(cycle: TreatmentCycle, assessment_name: str = "PHQ-9") -> str:
    if not cycle.entries:
        return f"{cycle.label} ({cycle.start_date} to {cycle.end_date}): no {assessment_name} check-ins were recorded."
    totals = [row.total for row in cycle.entries]
    return (
        f"{cycle.label} ({cycle.start_date} to {cycle.end_date}): {len(totals)} recorded {assessment_name} check-ins; "
        f"scores ranged from {min(totals)} to {max(totals)} with an average of {sum(totals) / len(totals):.1f}."
    )


def parse_date(value) -> str | None:
    if value is None or value == "":
        return None
    if hasattr(value, "date"):
        return value.date().isoformat()
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return None


def init_db(db_path: Path | None = None) -> None:
    db_path = db_path or DB_PATH
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(db_path)) as conn, conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS phq9_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                entry_date TEXT NOT NULL UNIQUE,
                item1 INTEGER NOT NULL,
                item2 INTEGER NOT NULL,
                item3 INTEGER NOT NULL,
                item4 INTEGER NOT NULL,
                item5 INTEGER NOT NULL,
                item6 INTEGER NOT NULL,
                item7 INTEGER NOT NULL,
                item8 INTEGER NOT NULL,
                item9 INTEGER NOT NULL,
                total INTEGER NOT NULL,
                severity TEXT NOT NULL,
                notes TEXT DEFAULT '',
                note_tag TEXT DEFAULT '',
                source TEXT DEFAULT 'manual',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        existing_columns = {row[1] for row in conn.execute("PRAGMA table_info(phq9_entries)").fetchall()}
        if "note_tag" not in existing_columns:
            conn.execute("ALTER TABLE phq9_entries ADD COLUMN note_tag TEXT DEFAULT ''")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS assessment_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                assessment_id TEXT NOT NULL,
                entry_date TEXT NOT NULL,
                item1 INTEGER NOT NULL,
                item2 INTEGER NOT NULL,
                item3 INTEGER NOT NULL,
                item4 INTEGER NOT NULL,
                item5 INTEGER NOT NULL,
                item6 INTEGER NOT NULL,
                item7 INTEGER NOT NULL,
                item8 INTEGER,
                item9 INTEGER,
                total INTEGER NOT NULL,
                severity TEXT NOT NULL,
                notes TEXT DEFAULT '',
                note_tag TEXT DEFAULT '',
                source TEXT DEFAULT 'manual',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(assessment_id, entry_date)
            )
            """
        )
        conn.execute(
            """
            INSERT OR IGNORE INTO assessment_entries (
                assessment_id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                total, severity, notes, note_tag, source, created_at, updated_at
            )
            SELECT 'phq9', entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                   total, severity, COALESCE(notes, ''), COALESCE(note_tag, ''), source, created_at, updated_at
            FROM phq9_entries
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS treatment_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_date TEXT NOT NULL,
                event_type TEXT NOT NULL,
                description TEXT DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        conn.commit()


def upsert_assessment_entry(
    assessment_id: str,
    entry_date: str,
    items: list[int],
    notes: str = "",
    source: str = "manual",
    note_tag: str = "",
) -> None:
    definition = ASSESSMENTS[assessment_id]
    if len(items) != definition.item_count:
        raise ValueError(f"{definition.display_name} requires {definition.item_count} items.")
    if any(score < 0 or score > 3 for score in items):
        raise ValueError(f"{definition.display_name} item scores must be 0, 1, 2, or 3.")
    padded_items = [*items, *([None] * (9 - len(items)))]
    total = sum(items)
    severity = severity_for_assessment(assessment_id, total)
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute(
            """
            INSERT INTO assessment_entries (
                assessment_id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                total, severity, notes, note_tag, source
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(assessment_id, entry_date) DO UPDATE SET
                item1=excluded.item1,
                item2=excluded.item2,
                item3=excluded.item3,
                item4=excluded.item4,
                item5=excluded.item5,
                item6=excluded.item6,
                item7=excluded.item7,
                item8=excluded.item8,
                item9=excluded.item9,
                total=excluded.total,
                severity=excluded.severity,
                notes=CASE
                    WHEN excluded.notes != '' THEN excluded.notes
                    ELSE assessment_entries.notes
                END,
                note_tag=CASE
                    WHEN excluded.note_tag != '' THEN excluded.note_tag
                    ELSE assessment_entries.note_tag
                END,
                source=excluded.source,
                updated_at=CURRENT_TIMESTAMP
            """,
            (assessment_id, entry_date, *padded_items, total, severity, notes or "", note_tag or "", source),
        )
        conn.commit()


def upsert_entry(entry_date: str, items: list[int], notes: str = "", source: str = "manual", note_tag: str = "") -> None:
    total = sum(items)
    severity = severity_for_score(total)
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute(
            """
            INSERT INTO phq9_entries (
                entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                total, severity, notes, note_tag, source
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(entry_date) DO UPDATE SET
                item1=excluded.item1,
                item2=excluded.item2,
                item3=excluded.item3,
                item4=excluded.item4,
                item5=excluded.item5,
                item6=excluded.item6,
                item7=excluded.item7,
                item8=excluded.item8,
                item9=excluded.item9,
                total=excluded.total,
                severity=excluded.severity,
                notes=CASE
                    WHEN excluded.notes != '' THEN excluded.notes
                    ELSE phq9_entries.notes
                END,
                note_tag=CASE
                    WHEN excluded.note_tag != '' THEN excluded.note_tag
                    ELSE phq9_entries.note_tag
                END,
                source=excluded.source,
                updated_at=CURRENT_TIMESTAMP
            """,
            (entry_date, *items, total, severity, notes or "", note_tag or "", source),
        )
        conn.commit()
    upsert_assessment_entry("phq9", entry_date, items, notes=notes, source=source, note_tag=note_tag)


def fetch_assessment_entries(
    assessment_id: str,
    start: str | None = None,
    end: str | None = None,
    limit: int | None = None,
) -> list[AssessmentEntryRow]:
    definition = ASSESSMENTS[assessment_id]
    clauses = ["assessment_id = ?"]
    params = [assessment_id]
    if start:
        clauses.append("entry_date >= ?")
        params.append(start)
    if end:
        clauses.append("entry_date <= ?")
        params.append(end)
    limit_sql = f"LIMIT {int(limit)}" if limit else ""
    sql = f"""
        SELECT id, assessment_id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
               total, severity, COALESCE(notes, ''), COALESCE(note_tag, '')
        FROM assessment_entries
        WHERE {" AND ".join(clauses)}
        ORDER BY entry_date ASC
        {limit_sql}
    """
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        rows = conn.execute(sql, params).fetchall()
    return [
        AssessmentEntryRow(
            id=row[0],
            assessment_id=row[1],
            entry_date=row[2],
            items=[int(v) for v in row[3 : 3 + definition.item_count]],
            total=int(row[12]),
            severity=row[13],
            notes=row[14] or "",
            note_tag=row[15] or "",
        )
        for row in rows
    ]


def fetch_entries(start: str | None = None, end: str | None = None, limit: int | None = None) -> list[EntryRow]:
    clauses = []
    params = []
    if start:
        clauses.append("entry_date >= ?")
        params.append(start)
    if end:
        clauses.append("entry_date <= ?")
        params.append(end)
    where = "WHERE " + " AND ".join(clauses) if clauses else ""
    limit_sql = f"LIMIT {int(limit)}" if limit else ""
    sql = f"""
        SELECT id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
               total, severity, COALESCE(notes, ''), COALESCE(note_tag, '')
        FROM phq9_entries
        {where}
        ORDER BY entry_date ASC
        {limit_sql}
    """
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        rows = conn.execute(sql, params).fetchall()
    return [
        EntryRow(
            id=row[0],
            entry_date=row[1],
            items=[int(v) for v in row[2:11]],
            total=int(row[11]),
            severity=row[12],
            notes=row[13] or "",
            note_tag=row[14] or "",
        )
        for row in rows
    ]


def fetch_recent_entries(days: int = 14) -> list[EntryRow]:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        rows = conn.execute(
            """
            SELECT id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                   total, severity, COALESCE(notes, ''), COALESCE(note_tag, '')
            FROM phq9_entries
            ORDER BY entry_date DESC
            LIMIT ?
            """,
            (days,),
        ).fetchall()
    rows.reverse()
    return [
        EntryRow(row[0], row[1], [int(v) for v in row[2:11]], int(row[11]), row[12], row[13] or "", row[14] or "")
        for row in rows
    ]


def fetch_events(start: str | None = None, end: str | None = None) -> list[tuple[int, str, str, str]]:
    clauses = []
    params = []
    if start:
        clauses.append("event_date >= ?")
        params.append(start)
    if end:
        clauses.append("event_date <= ?")
        params.append(end)
    where = "WHERE " + " AND ".join(clauses) if clauses else ""
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        return conn.execute(
            f"""
            SELECT id, event_date, event_type, COALESCE(description, '')
            FROM treatment_events
            {where}
            ORDER BY event_date ASC, id ASC
            """,
            params,
        ).fetchall()


def fetch_day_data(entry_date: str) -> dict[str, object]:
    """Return all editable records for one calendar date."""
    assessments = {
        assessment_id: rows[0]
        for assessment_id in ASSESSMENT_ORDER
        if (rows := fetch_assessment_entries(assessment_id, entry_date, entry_date))
    }
    notes = ""
    note_tag = ""
    for assessment_id in ASSESSMENT_ORDER:
        row = assessments.get(assessment_id)
        if row and (row.notes or row.note_tag):
            notes = row.notes
            note_tag = row.note_tag
            break
    return {
        "assessments": assessments,
        "notes": notes,
        "note_tag": note_tag,
        "events": fetch_events(entry_date, entry_date),
    }


def update_assessment_entry(entry_id: int, assessment_id: str, items: list[int]) -> None:
    """Update an assessment in place so its stable record ID is preserved."""
    definition = ASSESSMENTS[assessment_id]
    if len(items) != definition.item_count or any(score < 0 or score > 3 for score in items):
        raise ValueError(f"{definition.display_name} requires {definition.item_count} item scores from 0 to 3.")
    padded_items = [*items, *([None] * (9 - len(items)))]
    total = sum(items)
    severity = severity_for_assessment(assessment_id, total)
    assignments = ", ".join([*(f"item{i} = ?" for i in range(1, 10)), "total = ?", "severity = ?", "updated_at = CURRENT_TIMESTAMP"])
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        cursor = conn.execute(
            f"UPDATE assessment_entries SET {assignments} WHERE id = ? AND assessment_id = ?",
            (*padded_items, total, severity, entry_id, assessment_id),
        )
        if cursor.rowcount != 1:
            raise ValueError("The selected assessment no longer exists.")
        if assessment_id == "phq9":
            entry_date_row = conn.execute("SELECT entry_date FROM assessment_entries WHERE id = ?", (entry_id,)).fetchone()
            if entry_date_row:
                legacy_assignments = ", ".join([*(f"item{i} = ?" for i in range(1, 10)), "total = ?", "severity = ?", "updated_at = CURRENT_TIMESTAMP"])
                conn.execute(
                    f"UPDATE phq9_entries SET {legacy_assignments} WHERE entry_date = ?",
                    (*items, total, severity, entry_date_row[0]),
                )


def update_daily_note(entry_date: str, notes: str, note_tag: str = "") -> None:
    """Synchronize the single day-level note across assessment and legacy rows."""
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute(
            "UPDATE assessment_entries SET notes = ?, note_tag = ?, updated_at = CURRENT_TIMESTAMP WHERE entry_date = ?",
            (notes, note_tag, entry_date),
        )
        conn.execute(
            "UPDATE phq9_entries SET notes = ?, note_tag = ?, updated_at = CURRENT_TIMESTAMP WHERE entry_date = ?",
            (notes, note_tag, entry_date),
        )


def delete_daily_note(entry_date: str) -> None:
    update_daily_note(entry_date, "", "")


def delete_assessment_entry(entry_id: int, assessment_id: str) -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        row = conn.execute(
            "SELECT entry_date FROM assessment_entries WHERE id = ? AND assessment_id = ?",
            (entry_id, assessment_id),
        ).fetchone()
        if not row:
            return
        conn.execute("DELETE FROM assessment_entries WHERE id = ? AND assessment_id = ?", (entry_id, assessment_id))
        if assessment_id == "phq9":
            conn.execute("DELETE FROM phq9_entries WHERE entry_date = ?", (row[0],))


def update_event(event_id: int, event_date: str, event_type: str, description: str = "") -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        cursor = conn.execute(
            "UPDATE treatment_events SET event_date = ?, event_type = ?, description = ? WHERE id = ?",
            (event_date, event_type, description, event_id),
        )
        if cursor.rowcount != 1:
            raise ValueError("The selected treatment event no longer exists.")


def delete_event(event_id: int) -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("DELETE FROM treatment_events WHERE id = ?", (event_id,))


def import_spreadsheet(path: str) -> int:
    if load_workbook is None:
        raise RuntimeError("The openpyxl package is required for Excel import.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {"date", *(f"item {i}" for i in range(1, 10))}
    target = None
    headers = None
    for sheet in workbook.worksheets:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            continue
        names = {str(c).strip().lower() for c in first_row if c is not None}
        if required.issubset(names):
            target = sheet
            headers = [str(c).strip() if c is not None else "" for c in first_row]
            break
    if target is None or headers is None:
        raise ValueError("No sheet was found with Date and Item 1 through Item 9 columns.")

    col_index = {name.lower(): idx for idx, name in enumerate(headers)}
    count = 0
    blank_streak = 0
    for row in target.iter_rows(min_row=2, values_only=True):
        raw_date = row[col_index["date"]] if col_index["date"] < len(row) else None
        entry_date = parse_date(raw_date)
        if not entry_date:
            blank_streak += 1
            if blank_streak >= 50:
                break
            continue
        blank_streak = 0
        items = []
        for i in range(1, 10):
            idx = col_index[f"item {i}"]
            value = row[idx] if idx < len(row) else None
            if value is None or str(value).strip().lower() == "nan":
                items = []
                break
            score = int(value)
            if score < 0 or score > 3:
                raise ValueError(f"Item {i} on {entry_date} has value {score}; expected 0-3.")
            items.append(score)
        if len(items) != 9:
            continue
        notes = ""
        if "notes" in col_index and col_index["notes"] < len(row) and row[col_index["notes"]]:
            notes = str(row[col_index["notes"]])
        note_tag = ""
        if "note tag" in col_index and col_index["note tag"] < len(row) and row[col_index["note tag"]]:
            note_tag = str(row[col_index["note tag"]])
        upsert_entry(
            entry_date,
            items,
            notes=notes if notes.lower() != "nan" else "",
            source="spreadsheet",
            note_tag=note_tag if note_tag.lower() != "nan" else "",
        )
        for event_name in ("Ketamine", "Therapy"):
            key = event_name.lower()
            if key in col_index:
                value = row[col_index[key]] if col_index[key] < len(row) else None
                if bool(value) and str(value).lower() not in ("nan", "false", "0"):
                    add_event(entry_date, event_name, f"{event_name} marked in imported spreadsheet", dedupe=True)
        count += 1
    workbook.close()
    return count


def add_event(event_date: str, event_type: str, description: str = "", dedupe: bool = True) -> int:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        if dedupe:
            found = conn.execute(
                """
                SELECT id FROM treatment_events
                WHERE event_date = ? AND event_type = ? AND description = ?
                """,
                (event_date, event_type, description),
            ).fetchone()
            if found:
                return int(found[0])
        cursor = conn.execute(
            "INSERT INTO treatment_events (event_date, event_type, description) VALUES (?, ?, ?)",
            (event_date, event_type, description),
        )
        conn.commit()
        return int(cursor.lastrowid)


def upsert_daily_event(event_date: str, event_type: str, description: str = "") -> int:
    """Update the day's existing event tag or insert it once.

    The explicit event-management workflow can still add another legitimate event.
    """
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        found = conn.execute(
            "SELECT id FROM treatment_events WHERE event_date = ? AND event_type = ? ORDER BY id LIMIT 1",
            (event_date, event_type),
        ).fetchone()
        if found:
            if description:
                conn.execute("UPDATE treatment_events SET description = ? WHERE id = ?", (description, found[0]))
            return int(found[0])
        cursor = conn.execute(
            "INSERT INTO treatment_events (event_date, event_type, description) VALUES (?, ?, ?)",
            (event_date, event_type, description),
        )
        return int(cursor.lastrowid)


def _analysis_workbook_data() -> dict[str, list[dict[str, object]]]:
    """Build normalized, analysis-ready records without changing the database schema."""
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        assessments = conn.execute(
            """
            SELECT id, assessment_id, entry_date, item1, item2, item3, item4, item5,
                   item6, item7, item8, item9, total, severity, COALESCE(notes, '') AS notes,
                   COALESCE(note_tag, '') AS note_tag, source, created_at, updated_at
            FROM assessment_entries
            ORDER BY entry_date, assessment_id, id
            """
        ).fetchall()
        events = conn.execute(
            """
            SELECT id, event_date, event_type, COALESCE(description, '') AS description, created_at
            FROM treatment_events
            ORDER BY event_date, id
            """
        ).fetchall()

    daily_assessments: list[dict[str, object]] = []
    item_responses: list[dict[str, object]] = []
    notes_by_date: dict[str, dict[str, object]] = {}
    assessment_ids_by_date: dict[str, list[str]] = {}
    assessment_summary_by_date: dict[str, dict[str, object]] = {}

    for row in assessments:
        definition = ASSESSMENTS[row["assessment_id"]]
        assessment_record_id = f"assessment:{row['id']}"
        daily_record_id = f"day:{row['entry_date']}"
        daily_assessments.append(
            {
                "assessment_record_id": assessment_record_id,
                "assessment_entry_id": row["id"],
                "daily_record_id": daily_record_id,
                "entry_date": row["entry_date"],
                "assessment_id": row["assessment_id"],
                "assessment_name": definition.display_name,
                "daily_severity_score": row["total"],
                "severity_category": row["severity"],
                "source": row["source"],
                "created_at": row["created_at"],
                "updated_at": row["updated_at"],
            }
        )
        assessment_ids_by_date.setdefault(row["entry_date"], []).append(assessment_record_id)
        assessment_summary_by_date.setdefault(row["entry_date"], {})[f"{row['assessment_id']}_daily_severity_score"] = row["total"]
        assessment_summary_by_date[row["entry_date"]][f"{row['assessment_id']}_severity_category"] = row["severity"]
        for item_number, item_label in enumerate(definition.item_labels, start=1):
            response_value = row[f"item{item_number}"]
            item_responses.append(
                {
                    "item_response_record_id": f"{assessment_record_id}:item:{item_number}",
                    "assessment_record_id": assessment_record_id,
                    "assessment_entry_id": row["id"],
                    "daily_record_id": daily_record_id,
                    "entry_date": row["entry_date"],
                    "assessment_id": row["assessment_id"],
                    "item_number": item_number,
                    "item_label": item_label,
                    "response_value": response_value,
                    "symptom_present": bool(response_value > 0),
                }
            )
        if (row["notes"] or row["note_tag"]) and row["entry_date"] not in notes_by_date:
            notes_by_date[row["entry_date"]] = {
                "note_record_id": f"note:{row['entry_date']}",
                "daily_record_id": daily_record_id,
                "entry_date": row["entry_date"],
                "note_tag": row["note_tag"],
                "note_text": row["notes"],
                "created_at": row["created_at"],
                "updated_at": row["updated_at"],
            }

    treatment_events = []
    event_ids_by_date: dict[str, list[str]] = {}
    event_types_by_date: dict[str, list[str]] = {}
    for row in events:
        event_record_id = f"event:{row['id']}"
        normalized_type = normalize_event_type(row["event_type"])
        treatment_events.append(
            {
                "treatment_event_record_id": event_record_id,
                "treatment_event_id": row["id"],
                "daily_record_id": f"day:{row['event_date']}",
                "event_date": row["event_date"],
                "event_type": row["event_type"],
                "normalized_event_type": normalized_type,
                "description": row["description"],
                "created_at": row["created_at"],
            }
        )
        event_ids_by_date.setdefault(row["event_date"], []).append(event_record_id)
        event_types_by_date.setdefault(row["event_date"], []).append(normalized_type)

    ketamine_anchors: list[tuple[str, str]] = []
    for row in treatment_events:
        if row["normalized_event_type"] == "Ketamine" and not any(anchor[0] == row["event_date"] for anchor in ketamine_anchors):
            ketamine_anchors.append((str(row["event_date"]), str(row["treatment_event_record_id"])))
    all_dates = sorted(set(assessment_ids_by_date) | set(notes_by_date) | set(event_ids_by_date))
    last_recorded_date = all_dates[-1] if all_dates else None
    treatment_cycle_rows = []
    for index, (anchor_date, anchor_event_id) in enumerate(ketamine_anchors):
        next_anchor_date = ketamine_anchors[index + 1][0] if index + 1 < len(ketamine_anchors) else None
        cycle_end = (
            (datetime.fromisoformat(next_anchor_date).date() - timedelta(days=1)).isoformat()
            if next_anchor_date
            else last_recorded_date or anchor_date
        )
        treatment_cycle_rows.append(
            {
                "treatment_cycle_record_id": f"cycle:{anchor_date}",
                "anchor_treatment_event_record_id": anchor_event_id,
                "cycle_number": index + 1,
                "cycle_start_date": anchor_date,
                "cycle_end_date": cycle_end,
                "is_current_cycle": index == len(ketamine_anchors) - 1,
            }
        )

    daily_summary = []
    for entry_date in all_dates:
        summary = assessment_summary_by_date.get(entry_date, {})
        event_types = event_types_by_date.get(entry_date, [])
        daily_summary.append(
            {
                "daily_record_id": f"day:{entry_date}",
                "entry_date": entry_date,
                "assessment_record_ids": "|".join(assessment_ids_by_date.get(entry_date, [])),
                "phq9_daily_severity_score": summary.get("phq9_daily_severity_score", ""),
                "phq9_severity_category": summary.get("phq9_severity_category", ""),
                "gad7_daily_severity_score": summary.get("gad7_daily_severity_score", ""),
                "gad7_severity_category": summary.get("gad7_severity_category", ""),
                "note_record_id": notes_by_date.get(entry_date, {}).get("note_record_id", ""),
                "treatment_event_record_ids": "|".join(event_ids_by_date.get(entry_date, [])),
                "treatment_event_count": len(event_ids_by_date.get(entry_date, [])),
                "ketamine_recorded": "Ketamine" in event_types,
                "therapy_recorded": "Therapy" in event_types,
                "medication_change_recorded": any(event_type.startswith("Medication") for event_type in event_types),
            }
        )

    metadata = [
        {"metadata_key": "workbook_schema_version", "metadata_value": ANALYSIS_WORKBOOK_SCHEMA_VERSION, "description": "Version of this normalized export layout."},
        {"metadata_key": "generated_at", "metadata_value": datetime.now().astimezone().isoformat(timespec="seconds"), "description": "Local generation timestamp in ISO 8601 format."},
        {"metadata_key": "application", "metadata_value": "Mental Health Tracker", "description": "Application that generated the workbook."},
        {"metadata_key": "date_format", "metadata_value": "YYYY-MM-DD", "description": "Calendar-date format used in all date fields."},
        {"metadata_key": "daily_relationship", "metadata_value": "daily_record_id", "description": "Join records across worksheets by the deterministic day:YYYY-MM-DD identifier."},
        {"metadata_key": "assessment_relationship", "metadata_value": "assessment_record_id", "description": "Join Item Responses to Daily Assessments by assessment_record_id."},
        {"metadata_key": "event_relationship", "metadata_value": "treatment_event_record_id", "description": "Treatment Events retain the stable SQLite event ID as event:<id>."},
        {"metadata_key": "cycle_relationship", "metadata_value": "anchor_treatment_event_record_id", "description": "Treatment Cycles reference the first recorded ketamine event on each anchor date."},
        {"metadata_key": "daily_severity_score", "metadata_value": "sum of item responses on one check-in", "description": "Measures recorded symptom severity for a single assessment date."},
        {"metadata_key": "14_day_frequency_score", "metadata_value": "derived by calendar window", "description": "Counts symptom-present days and is intentionally not repeated as a daily raw field."},
        {"metadata_key": "missing_checkins", "metadata_value": "missing information", "description": "A day without a check-in is not evidence that symptoms were absent."},
        {"metadata_key": "daily_summary_authority", "metadata_value": "derived", "description": "Daily Summary is a convenience view; normalized worksheets remain authoritative."},
        {"metadata_key": "privacy_notice", "metadata_value": "local sensitive data", "description": "This workbook may contain PHI. Store and share it deliberately."},
    ]

    return {
        "Daily Assessments": daily_assessments,
        "Item Responses": item_responses,
        "Notes": list(notes_by_date.values()),
        "Treatment Events": treatment_events,
        "Treatment Cycles": treatment_cycle_rows,
        "Metadata": metadata,
        "Daily Summary": daily_summary,
    }


def export_analysis_workbook(path: str) -> None:
    """Write the separate normalized XLSX export used for external analysis."""
    if not path.lower().endswith(".xlsx"):
        raise ValueError("The analysis-ready export must be saved as an .xlsx workbook.")
    if pd is None:
        raise RuntimeError("Analysis-ready Excel export requires pandas/openpyxl.")
    sheet_columns = {
        "Daily Assessments": ["assessment_record_id", "assessment_entry_id", "daily_record_id", "entry_date", "assessment_id", "assessment_name", "daily_severity_score", "severity_category", "source", "created_at", "updated_at"],
        "Item Responses": ["item_response_record_id", "assessment_record_id", "assessment_entry_id", "daily_record_id", "entry_date", "assessment_id", "item_number", "item_label", "response_value", "symptom_present"],
        "Notes": ["note_record_id", "daily_record_id", "entry_date", "note_tag", "note_text", "created_at", "updated_at"],
        "Treatment Events": ["treatment_event_record_id", "treatment_event_id", "daily_record_id", "event_date", "event_type", "normalized_event_type", "description", "created_at"],
        "Treatment Cycles": ["treatment_cycle_record_id", "anchor_treatment_event_record_id", "cycle_number", "cycle_start_date", "cycle_end_date", "is_current_cycle"],
        "Metadata": ["metadata_key", "metadata_value", "description"],
        "Daily Summary": ["daily_record_id", "entry_date", "assessment_record_ids", "phq9_daily_severity_score", "phq9_severity_category", "gad7_daily_severity_score", "gad7_severity_category", "note_record_id", "treatment_event_record_ids", "treatment_event_count", "ketamine_recorded", "therapy_recorded", "medication_change_recorded"],
    }
    workbook_data = _analysis_workbook_data()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, columns in sheet_columns.items():
            pd.DataFrame(workbook_data[sheet_name], columns=columns).to_excel(writer, index=False, sheet_name=sheet_name)


def python_supports_modules(python_path: Path, modules: tuple[str, ...]) -> bool:
    if not python_path.exists():
        return False
    imports = "; ".join(f"import {module}" for module in modules)
    try:
        result = subprocess.run(
            [str(python_path), "-c", imports],
            capture_output=True,
            text=True,
            timeout=15,
        )
        return result.returncode == 0
    except (OSError, subprocess.SubprocessError):
        return False


def find_helper_python(required_modules: tuple[str, ...]) -> Path | None:
    """Find a configured or project-local Python that has the requested features."""
    candidates = []
    if BUNDLED_PYTHON is not None:
        candidates.append(BUNDLED_PYTHON)
    candidates.extend(
        [
            PROJECT_ROOT / ".venv" / "Scripts" / "python.exe",
            PROJECT_ROOT / "venv" / "Scripts" / "python.exe",
        ]
    )
    current = Path(sys.executable).resolve()
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if resolved == current:
            continue
        if python_supports_modules(resolved, required_modules):
            return resolved
    return None


def run_bundled_cli(args: list[str], required_modules: tuple[str, ...]) -> None:
    helper_python = find_helper_python(required_modules)
    if helper_python is None:
        names = ", ".join(required_modules)
        raise RuntimeError(
            f"This action needs Python packages that are not available: {names}. "
            "Install the project requirements in .venv, venv, or the active Python environment, "
            "or set PHQ9_TRACKER_BUNDLED_PYTHON to a compatible python.exe."
        )
    command = [str(helper_python), str(Path(__file__).resolve()), *args]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "Bundled command failed.")


def add_pdf_table(
    story,
    title: str,
    rows: list[list[str]],
    repeat_header: bool = True,
    highlight_flag_col: int | None = None,
    col_widths: list[float] | None = None,
    wrap_columns: set[int] | None = None,
) -> None:
    styles = getSampleStyleSheet()
    story.append(Paragraph(title, styles["Heading2"]))
    if wrap_columns:
        wrapped = []
        for row_idx, row in enumerate(rows):
            wrapped_row = []
            for col_idx, value in enumerate(row):
                if row_idx > 0 and col_idx in wrap_columns:
                    safe_value = escape(str(value)).replace("\n", "<br/>")
                    wrapped_row.append(Paragraph(safe_value, styles["BodyText"]))
                else:
                    wrapped_row.append(str(value))
            wrapped.append(wrapped_row)
        rows = wrapped
    table = Table(rows, repeatRows=1 if repeat_header and len(rows) > 1 else 0, colWidths=col_widths)
    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#172033")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]
    if highlight_flag_col is not None:
        for row_idx, row in enumerate(rows[1:], start=1):
            if highlight_flag_col < len(row) and str(row[highlight_flag_col]).strip():
                style_commands.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#FFF7ED")))
                style_commands.append(("TEXTCOLOR", (highlight_flag_col, row_idx), (highlight_flag_col, row_idx), colors.HexColor("#9A3412")))
    table.setStyle(TableStyle(style_commands)
    )
    story.append(table)
    story.append(Spacer(1, 0.14 * inch))


def normalize_event_type(event_type: str) -> str:
    text = (event_type or "").strip().lower()
    if "ketamine" in text:
        return "Ketamine"
    if "therapy" in text:
        return "Therapy"
    if "dose increase" in text:
        return "Medication Dose Increase"
    if "dose decrease" in text:
        return "Medication Dose Decrease"
    if "start" in text:
        return "Medication Start"
    if "stop" in text or "discontinued" in text:
        return "Medication Stop"
    return event_type or "Treatment event"


def draw_line_chart(
    path: str,
    entries: list[EntryRow],
    series: list[tuple[str, list[float | int | None], str]],
    title: str,
    y_max: int,
    events: list[tuple[int, str, str, str]] | None = None,
    show_severity: bool = False,
    note_markers: bool = False,
    width: int = 1100,
    height: int = 520,
) -> None:
    if PILImage is None:
        raise RuntimeError("Chart generation requires Pillow.")
    img = PILImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    title_font = ImageFont.load_default()
    left, right, top, bottom = 70, 35, 58, 72
    plot_w = width - left - right
    plot_h = height - top - bottom
    draw.text((left, 18), title, fill="#172033", font=title_font)
    if not entries:
        draw.text((width // 2 - 60, height // 2), "No entries", fill="#64748B", font=font)
        img.save(path)
        return

    def x_for_idx(idx: int) -> float:
        return left + (idx / max(len(entries) - 1, 1)) * plot_w

    def y_for_value(value: float) -> float:
        return top + plot_h - (value / y_max) * plot_h

    if show_severity:
        bands = [
            (0, 4, "#ECFDF5", "Minimal"),
            (5, 9, "#F0FDF4", "Mild"),
            (10, 14, "#FEFCE8", "Moderate"),
            (15, 19, "#FFF7ED", "Moderately Severe"),
            (20, 27, "#FEF2F2", "Severe"),
        ]
        for low, high, color, label in bands:
            y1 = y_for_value(high)
            y2 = y_for_value(low)
            draw.rectangle([left, y1, left + plot_w, y2], fill=color)
            draw.text((left + 6, y1 + 4), label, fill="#475569", font=font)

    for tick in range(0, y_max + 1, 3 if y_max > 5 else 1):
        y = y_for_value(tick)
        draw.line([(left, y), (left + plot_w, y)], fill="#E2E8F0")
        draw.text((18, y - 6), str(tick), fill="#475569", font=font)
    draw.line([(left, top), (left, top + plot_h), (left + plot_w, top + plot_h)], fill="#94A3B8", width=2)

    event_lookup = {}
    for event in events or []:
        event_lookup.setdefault(event[1], []).append(event)
    for idx, row in enumerate(entries):
        x = x_for_idx(idx)
        for event in event_lookup.get(row.entry_date, []):
            kind = normalize_event_type(event[2])
            color = {
                "Ketamine": "#7C3AED",
                "Therapy": "#0F766E",
                "Medication Start": "#2563EB",
                "Medication Stop": "#DC2626",
                "Medication Dose Increase": "#EA580C",
                "Medication Dose Decrease": "#0891B2",
            }.get(kind, "#64748B")
            draw.line([(x, top), (x, top + plot_h)], fill=color, width=2)
            draw.polygon([(x, top - 8), (x - 5, top), (x + 5, top)], fill=color)
        if note_markers and row.notes.strip():
            y = y_for_value(row.total)
            draw.ellipse([x - 5, y - 5, x + 5, y + 5], outline="#111827", width=2)

    for name, values, color in series:
        points = []
        for idx, value in enumerate(values):
            if value is None:
                if len(points) > 1:
                    draw.line(points, fill=color, width=3)
                points = []
                continue
            points.append((x_for_idx(idx), y_for_value(float(value))))
        if len(points) > 1:
            draw.line(points, fill=color, width=3)
        for x, y in points[:: max(1, len(points) // 35)]:
            draw.ellipse([x - 3, y - 3, x + 3, y + 3], fill=color)

    draw.text((left, height - 44), entries[0].entry_date, fill="#475569", font=font)
    draw.text((width - 110, height - 44), entries[-1].entry_date, fill="#475569", font=font)
    legend_x, legend_y = left, height - 24
    for name, _, color in series:
        draw.rectangle([legend_x, legend_y, legend_x + 12, legend_y + 12], fill=color)
        draw.text((legend_x + 18, legend_y - 1), name, fill="#334155", font=font)
        legend_x += 150
    if events:
        draw.text((legend_x, legend_y - 1), "Vertical markers: treatment/medication events", fill="#64748B", font=font)
    if note_markers:
        draw.text((left + 420, 18), "Outlined dots: notes", fill="#111827", font=font)
    img.save(path)


def add_chart(story, image_path: str, caption: str | None = None, width: float = 7.2 * inch) -> None:
    if caption:
        story.append(Paragraph(caption, getSampleStyleSheet()["Heading2"]))
    story.append(Image(image_path, width=width, height=width * 0.47))
    story.append(Spacer(1, 0.14 * inch))


def draw_treatment_bar_chart(path: str, title: str, rows: list[list[str]], labels: list[str]) -> None:
    if PILImage is None:
        raise RuntimeError("Chart generation requires Pillow.")
    width, height = 950, 360
    img = PILImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    draw.text((50, 18), title, fill="#172033", font=font)
    if not rows:
        draw.text((width // 2 - 80, height // 2), "No events recorded", fill="#64748B", font=font)
        img.save(path)
        return
    values_by_label = []
    for col_idx in range(2, 2 + len(labels)):
        values = []
        for row in rows:
            try:
                values.append(float(row[col_idx]))
            except (ValueError, IndexError):
                pass
        values_by_label.append(sum(values) / len(values) if values else None)
    left, top, bottom = 85, 68, 58
    plot_h = height - top - bottom
    plot_w = width - left - 45
    y_max = max([27, *[v for v in values_by_label if v is not None]])
    for tick in range(0, 28, 3):
        y = top + plot_h - (tick / y_max) * plot_h
        draw.line([(left, y), (left + plot_w, y)], fill="#E2E8F0")
        draw.text((35, y - 6), str(tick), fill="#475569", font=font)
    colors_local = ["#2563EB", "#7C3AED", "#0F766E"]
    bar_gap = plot_w / max(len(labels), 1)
    for idx, label in enumerate(labels):
        value = values_by_label[idx]
        x1 = left + idx * bar_gap + bar_gap * 0.2
        x2 = left + (idx + 1) * bar_gap - bar_gap * 0.2
        if value is not None:
            y = top + plot_h - (value / y_max) * plot_h
            draw.rectangle([x1, y, x2, top + plot_h], fill=colors_local[idx % len(colors_local)])
            draw.text((x1, y - 18), f"{value:.1f}", fill="#172033", font=font)
        draw.text((x1, top + plot_h + 12), label, fill="#334155", font=font)
    draw.line([(left, top), (left, top + plot_h), (left + plot_w, top + plot_h)], fill="#94A3B8", width=2)
    img.save(path)


def on_report_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(36, 20, DISCLAIMER)
    canvas.drawRightString(576, 20, f"Page {doc.page}")
    canvas.restoreState()


def available_report_date_range() -> tuple[str, str]:
    """Return the full assessment-history range used by one-click PDF reports."""
    dates = [
        row.entry_date
        for assessment_id in ASSESSMENT_ORDER
        for row in fetch_assessment_entries(assessment_id)
    ]
    if not dates:
        raise ValueError("No check-ins are available for a clinician report.")
    return min(dates), max(dates)


def generate_report(start: str, end: str, pdf_path: str) -> None:
    if colors is None or PILImage is None:
        raise RuntimeError("PDF export requires reportlab and Pillow.")
    entries = fetch_assessment_entries("phq9", start, end)
    gad_entries = fetch_assessment_entries("gad7", start, end)
    if not entries and not gad_entries:
        raise ValueError("No entries found in the selected date range.")
    comparison_start = (datetime.fromisoformat(end).date() - timedelta(days=27)).isoformat()
    comparison_phq_entries = fetch_assessment_entries("phq9", comparison_start, end)
    comparison_gad_entries = fetch_assessment_entries("gad7", comparison_start, end)
    events = [
        (event_id, event_date, normalize_event_type(event_type), description)
        for event_id, event_date, event_type, description in fetch_events(start, end)
    ]
    cycles = treatment_cycles(entries, events, end)
    highlights = [
        *symptom_highlights("phq9", comparison_phq_entries, end, limit=2),
        *symptom_highlights("gad7", comparison_gad_entries, end, limit=2),
    ]

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    chart_dir = Path(tempfile.mkdtemp(prefix="phq9_report_charts_"))
    phq_chart = chart_dir / "phq9_recent.png"
    gad_chart = chart_dir / "gad7_recent.png"
    draw_line_chart(
        str(phq_chart),
        entries[-90:],
        [("PHQ-9", [row.total for row in entries[-90:]], "#2563EB")],
        "PHQ-9 Recorded Score Trend",
        27,
        events=events,
        width=1100,
        height=360,
    )
    draw_line_chart(
        str(gad_chart),
        gad_entries[-90:],
        [("GAD-7", [row.total for row in gad_entries[-90:]], "#0F766E")],
        "GAD-7 Recorded Score Trend",
        21,
        events=events,
        width=1100,
        height=360,
    )

    story = [
        Paragraph("Mental Health Tracker Clinician Discussion Report", styles["Title"]),
        Paragraph(f"Date range: {start} to {end}", styles["Normal"]),
        Paragraph(DISCLAIMER, styles["BodyText"]),
        Spacer(1, 0.18 * inch),
        Paragraph("How to Read This Report", styles["Heading1"]),
        Paragraph(
            "<b>Daily Severity Score:</b> the item responses from one PHQ-9 or GAD-7 check-in are added together. "
            "It describes the severity recorded on that particular day.",
            styles["BodyText"],
        ),
        Paragraph(
            "<b>14-Day Symptom Frequency Score:</b> for each item, the report counts how many of the 14 calendar days "
            "had a recorded response above zero. A response of 1 and a response of 3 each count as one symptom-present "
            "day, even though they contribute differently to the Daily Severity Score. Item counts are converted to "
            "0 points for 0 days, 1 point for 1-6 days, 2 points for 7-11 days, or 3 points for 12-14 days, then summed.",
            styles["BodyText"],
        ),
        Paragraph(
            "<b>Coverage and missing check-ins:</b> coverage shows how many daily check-ins were recorded in a calendar "
            "window. A day without a check-in contributes no recorded symptom-present response to the frequency score, "
            "but it is missing information, not evidence that the symptom was absent. Interpret comparisons alongside coverage.",
            styles["BodyText"],
        ),
        Spacer(1, 0.12 * inch),
        Paragraph("Recorded Period Overview", styles["Heading1"]),
    ]
    glance_rows = [["Assessment", "Recorded check-ins", "Most recent score", "Most recent date"]]
    for assessment_id, assessment_entries in (("phq9", entries), ("gad7", gad_entries)):
        name = ASSESSMENTS[assessment_id].display_name
        glance_rows.append(
            [
                name,
                str(len(assessment_entries)),
                f"{assessment_entries[-1].total} ({assessment_entries[-1].severity})" if assessment_entries else "n/a",
                assessment_entries[-1].entry_date if assessment_entries else "n/a",
            ]
        )
    add_pdf_table(
        story,
        "Recorded Check-Ins",
        glance_rows,
        col_widths=[1.2 * inch, 1.35 * inch, 2.1 * inch, 1.45 * inch],
        wrap_columns={0, 1, 2, 3},
    )
    story.append(Paragraph("Overall pattern", styles["Heading2"]))
    story.append(Paragraph(overall_pattern_summary(comparison_phq_entries, comparison_gad_entries, end), styles["BodyText"]))
    story.append(Paragraph("Symptom highlights", styles["Heading2"]))
    for highlight in highlights:
        story.append(Paragraph(f"- {highlight}", styles["BodyText"]))
    story.append(Spacer(1, 0.08 * inch))
    story.append(
        Paragraph(
            "Highlights describe recorded check-ins only. A day without a check-in is missing information, not evidence that a symptom was absent.",
            styles["BodyText"],
        )
    )

    story.append(PageBreak())
    story.append(Paragraph("Recorded Symptom Trends", styles["Heading1"]))
    add_chart(story, str(phq_chart), "PHQ-9 scores across the selected period", width=6.1 * inch)
    add_chart(story, str(gad_chart), "GAD-7 scores across the selected period", width=6.1 * inch)
    q9_context = item9_context(entries, end)
    q9_summary = item9_context_summary(q9_context)
    if q9_summary:
        story.append(Paragraph("PHQ-9 item 9 context", styles["Heading2"]))
        story.append(
            Paragraph(
                q9_summary,
                styles["BodyText"],
            )
        )

    story.append(PageBreak())
    story.append(Paragraph("Treatment Context", styles["Heading1"]))
    story.append(
        Paragraph(
            "Cycles are bounded by recorded ketamine infusion dates. These observations describe timing and recorded scores; they do not determine whether treatment caused a change.",
            styles["BodyText"],
        )
    )
    if cycles:
        cycle_rows = [["Cycle", "Dates", "Recorded PHQ-9 pattern"]]
        for cycle in cycles:
            cycle_rows.append([cycle.label, f"{cycle.start_date} to {cycle.end_date}", treatment_cycle_observation(cycle)])
        add_pdf_table(
            story,
            "Current and Previous Recorded Cycles",
            cycle_rows,
            col_widths=[1.0 * inch, 1.65 * inch, 4.05 * inch],
            wrap_columns={0, 1, 2},
        )
        for cycle in cycles:
            cycle_chart = chart_dir / f"{cycle.label.lower().replace(' ', '_')}.png"
            draw_line_chart(
                str(cycle_chart),
                cycle.entries,
                [("PHQ-9", [row.total for row in cycle.entries], "#7C3AED")],
                f"{cycle.label}: Recorded PHQ-9 Scores",
                27,
                width=1100,
                height=300,
            )
            add_chart(story, str(cycle_chart), width=6.2 * inch)
    else:
        story.append(Paragraph("No ketamine infusion was recorded in the selected period, so cycle views are not available.", styles["BodyText"]))

    timeline = []
    for _, event_date, event_type, description in events:
        timeline.append((event_date, event_type, description or "Recorded event"))
    seen_notes = set()
    for row in [*entries, *gad_entries]:
        note_key = (row.entry_date, row.notes.strip(), row.note_tag.strip())
        if row.notes.strip() and note_key not in seen_notes:
            seen_notes.add(note_key)
            timeline.append((row.entry_date, f"Note{f' - {row.note_tag}' if row.note_tag else ''}", row.notes.strip()))
    timeline = sorted(timeline, key=lambda item: (item[0], item[1]))

    story.append(PageBreak())
    story.append(Paragraph("Journal and Event Context", styles["Heading1"]))
    if timeline:
        add_pdf_table(
            story,
            "Recorded Context",
            [["Date", "Type", "Recorded context"]]
            + [[event_date, event_type, text] for event_date, event_type, text in timeline],
            col_widths=[0.95 * inch, 1.45 * inch, 4.3 * inch],
            wrap_columns={1, 2},
        )
    else:
        story.append(Paragraph("No notes or treatment events were recorded in the selected period.", styles["BodyText"]))
    story.append(Paragraph("Conversation Starters", styles["Heading2"]))
    prompts = [
        "Do the recorded symptom patterns match what you remember about this period?",
        "Were there particular days, events, or treatment dates that would help explain the recorded context?",
        "Which symptom changes would be most useful to discuss or monitor together next?",
    ]
    q9_prompt = item9_conversation_starter(q9_context)
    if q9_prompt:
        prompts.insert(0, q9_prompt)
    for prompt in prompts[:4]:
        story.append(Paragraph(f"- {prompt}", styles["BodyText"]))
    doc.build(story, onFirstPage=on_report_page, onLaterPages=on_report_page)


class LineChart(Canvas):
    def __init__(self, parent, **kwargs):
        kwargs.setdefault("takefocus", True)
        super().__init__(
            parent,
            bg="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#CBD5E1",
            highlightcolor="#2563EB",
            **kwargs,
        )
        self._series_args = None
        self._chart_points: list[ChartPoint] = []
        self._hovered_point_index: int | None = None
        self._locked_point_index: int | None = None
        self.bind("<Configure>", self._redraw_after_resize)
        self.bind("<Motion>", self._show_hover_callout)
        self.bind("<Leave>", self._hide_hover_callout)
        self.bind("<Button-1>", self._toggle_click_callout)
        self.bind("<Left>", lambda event: self._move_keyboard_callout(-1))
        self.bind("<Right>", lambda event: self._move_keyboard_callout(1))
        self.bind("<Return>", self._lock_keyboard_callout)
        self.bind("<Escape>", self._clear_callout)

    def draw_series(self, entries: list[EntryRow], series: list[tuple[str, list[int], str]], y_max: int, title: str) -> None:
        self._series_args = (entries, series, y_max, title)
        self._render_series(entries, series, y_max, title)

    def _redraw_after_resize(self, _event=None) -> None:
        if self._series_args is not None:
            self._render_series(*self._series_args)

    def _render_series(self, entries: list[EntryRow], series: list[tuple[str, list[int], str]], y_max: int, title: str) -> None:
        locked_index = self._locked_point_index
        self.delete("all")
        self._chart_points = []
        self._hovered_point_index = None
        width = responsive_canvas_dimension(self.winfo_width(), int(self["width"]))
        height = responsive_canvas_dimension(self.winfo_height(), int(self["height"]))
        pad_l, pad_r, pad_b = 46, 18, 36
        title_id = self.create_text(
            12,
            10,
            text=title,
            width=max(120, width - 24),
            anchor="nw",
            justify=LEFT,
            fill="#172033",
            font=("Segoe UI", 11, "bold"),
            tags=("chart-title",),
        )
        title_bounds = self.bbox(title_id) or (12, 10, width - 12, 28)
        legend_top, pad_t = chart_vertical_positions(title_bounds[3])
        if not entries or not series:
            self.create_text(width / 2, max(pad_t + 20, height / 2), text="No entries to display", fill="#64748B")
            return
        plot_w = width - pad_l - pad_r
        plot_h = height - pad_t - pad_b
        self.create_line(pad_l, pad_t, pad_l, pad_t + plot_h, fill="#94A3B8")
        self.create_line(pad_l, pad_t + plot_h, pad_l + plot_w, pad_t + plot_h, fill="#94A3B8")
        for tick in range(0, y_max + 1, max(1, y_max // 5)):
            y = pad_t + plot_h - (tick / y_max) * plot_h
            self.create_line(pad_l - 4, y, pad_l + plot_w, y, fill="#E2E8F0")
            self.create_text(pad_l - 8, y, text=str(tick), anchor="e", fill="#475569", font=("Segoe UI", 8))
        points_count = max(len(entries) - 1, 1)
        for name, values, color in series:
            line_points = []
            series_points = []
            for idx, value in enumerate(values[: len(entries)]):
                x = pad_l + (idx / points_count) * plot_w
                y = pad_t + plot_h - (value / y_max) * plot_h
                line_points.extend([x, y])
                series_points.append(ChartPoint(x, y, entries[idx].entry_date, name, value, color))
            if len(line_points) >= 4:
                self.create_line(*line_points, fill=color, width=2, smooth=True, tags=("chart-line",))
            for point in series_points:
                self._chart_points.append(point)
                self.create_oval(
                    point.x - 4,
                    point.y - 4,
                    point.x + 4,
                    point.y + 4,
                    fill=point.color,
                    outline="#FFFFFF",
                    width=1,
                    tags=("chart-point",),
                )
        labels = [entries[0].entry_date, entries[-1].entry_date] if len(entries) > 1 else [entries[0].entry_date]
        self.create_text(pad_l, height - 14, text=labels[0], anchor="w", fill="#475569", font=("Segoe UI", 8))
        if len(labels) > 1:
            self.create_text(width - pad_r, height - 14, text=labels[1], anchor="e", fill="#475569", font=("Segoe UI", 8))
        legend_x = pad_l + 8
        for name, _, color in series:
            self.create_rectangle(legend_x, legend_top, legend_x + 10, legend_top + 10, fill=color, outline=color, tags=("chart-legend",))
            self.create_text(legend_x + 14, legend_top + 5, text=name, anchor="w", fill="#334155", font=("Segoe UI", 8), tags=("chart-legend",))
            legend_x += max(90, len(name) * 7)

        if locked_index is not None and locked_index < len(self._chart_points):
            self._locked_point_index = locked_index
            self._display_callout(locked_index)
        else:
            self._locked_point_index = None

    def _point_index_at(self, x: float, y: float) -> int | None:
        point = nearest_chart_point(self._chart_points, x, y)
        if point is None:
            return None
        return self._chart_points.index(point)

    def _display_callout(self, point_index: int) -> None:
        self.delete("chart-callout")
        point = self._chart_points[point_index]
        width = responsive_canvas_dimension(self.winfo_width(), int(self["width"]))
        text_x = point.x + 10 if point.x <= width * 0.62 else point.x - 10
        anchor = "sw" if point.x <= width * 0.62 else "se"
        text_y = point.y - 9
        if text_y < 28:
            text_y = point.y + 9
            anchor = "nw" if point.x <= width * 0.62 else "ne"
        text_id = self.create_text(
            text_x,
            text_y,
            text=chart_point_callout_text(point),
            anchor=anchor,
            fill="#172033",
            font=("Segoe UI", 9, "bold"),
            tags=("chart-callout",),
        )
        bounds = self.bbox(text_id)
        if bounds:
            rectangle_id = self.create_rectangle(
                bounds[0] - 6,
                bounds[1] - 4,
                bounds[2] + 6,
                bounds[3] + 4,
                fill="#FFF7D6",
                outline="#A16207",
                width=1,
                tags=("chart-callout",),
            )
            self.tag_lower(rectangle_id, text_id)
        self.create_oval(
            point.x - 7,
            point.y - 7,
            point.x + 7,
            point.y + 7,
            outline="#172033",
            width=2,
            tags=("chart-callout",),
        )

    def _show_hover_callout(self, event) -> None:
        if self._locked_point_index is not None:
            return
        point_index = self._point_index_at(event.x, event.y)
        self.config(cursor="hand2" if point_index is not None else "")
        if point_index is None:
            self._hovered_point_index = None
            self.delete("chart-callout")
            return
        if point_index != self._hovered_point_index:
            self._hovered_point_index = point_index
            self._display_callout(point_index)

    def _hide_hover_callout(self, _event=None) -> None:
        self.config(cursor="")
        if self._locked_point_index is None:
            self._hovered_point_index = None
            self.delete("chart-callout")

    def _toggle_click_callout(self, event) -> str:
        self.focus_set()
        point_index = self._point_index_at(event.x, event.y)
        if point_index is None:
            self._locked_point_index = None
            self._hovered_point_index = None
            self.delete("chart-callout")
            return "break"
        self._locked_point_index = None if point_index == self._locked_point_index else point_index
        self._hovered_point_index = point_index
        if self._locked_point_index is None:
            self.delete("chart-callout")
        else:
            self._display_callout(point_index)
        return "break"

    def _move_keyboard_callout(self, direction: int) -> str:
        if not self._chart_points:
            return "break"
        if self._locked_point_index is not None:
            current = self._locked_point_index
        elif self._hovered_point_index is not None:
            current = self._hovered_point_index
        else:
            current = 0 if direction < 0 else -1
        self._locked_point_index = (current + direction) % len(self._chart_points)
        self._hovered_point_index = self._locked_point_index
        self._display_callout(self._locked_point_index)
        return "break"

    def _lock_keyboard_callout(self, _event=None) -> str:
        if self._chart_points and self._locked_point_index is None:
            self._locked_point_index = self._hovered_point_index or 0
            self._display_callout(self._locked_point_index)
        return "break"

    def _clear_callout(self, _event=None) -> str:
        self._locked_point_index = None
        self._hovered_point_index = None
        self.delete("chart-callout")
        return "break"


PRIMARY_TAB_ORDER = (
    "Today's Check-In",
    "Review",
    "History / Manage Entries",
    "Treatment Events",
    "How Scoring Works",
)

REVIEW_ACTION_LABELS = (
    "Refresh",
    "Import Spreadsheet",
    "Generate PDF",
    "Analysis Workbook",
)

CHART_INTERACTION_HELP = (
    "Chart values: hover or click a point for its date and exact score. "
    "Keyboard: Tab to a chart, use Left/Right to move, Enter to show, and Escape to clear."
)


class PHQ9App(Tk):
    def __init__(self):
        super().__init__()
        self.title("Mental Health Tracker")
        self.geometry("1180x780")
        self.minsize(980, 680)
        self.configure(bg="#F8FAFC")
        try:
            self.iconbitmap(default=str(ICON_PATH))
        except Exception:
            pass
        init_db()
        self._checkin_snapshot = None
        self._loaded_checkin_date = None
        self._history_snapshot = None
        self._loaded_history_date = None
        self.configure_ui_styles()
        self.create_menu()
        self.create_widgets()
        self.load_checkin_date(confirm_unsaved=False)
        self.load_history_date(confirm_unsaved=False)
        self.refresh_all()

    def configure_ui_styles(self):
        """Apply restrained, readable defaults without introducing a theme dependency."""
        self.option_add("*Font", ("Segoe UI", 10))
        self.option_add("*Button.padX", 10)
        self.option_add("*Button.padY", 5)
        self.option_add("*Text.Font", ("Segoe UI", 10))
        style = ttk.Style(self)
        style.configure(".", font=("Segoe UI", 10))
        style.configure("TNotebook.Tab", padding=(12, 7))
        style.configure("Treeview", rowheight=26)
        style.configure("TSpinbox", padding=3)
        style.configure("TCombobox", padding=3)

    def create_menu(self):
        menu = Menu(self)
        file_menu = Menu(menu, tearoff=0)
        file_menu.add_command(label="Import spreadsheet...", command=self.import_file)
        file_menu.add_command(label="Export analysis workbook...", command=self.export_analysis_file)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.destroy)
        menu.add_cascade(label="File", menu=file_menu)
        help_menu = Menu(menu, tearoff=0)
        help_menu.add_command(label="How Scoring Works", command=self.show_scoring_help)
        menu.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menu)

    def create_widgets(self):
        header = Frame(self, bg="#172033", padx=16, pady=12)
        header.pack(fill="x")
        Label(header, text="Mental Health Tracker", fg="white", bg="#172033", font=("Segoe UI", 18, "bold")).pack(side=LEFT)
        Label(
            header,
            text="All data is stored locally in SQLite. No upload or cloud sync is performed by this app.",
            fg="#DDE7F3",
            bg="#172033",
            font=("Segoe UI", 10),
        ).pack(side=RIGHT)

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill=BOTH, expand=True, padx=12, pady=12)
        self.dashboard = Frame(self.notebook, bg="#F8FAFC")
        self.entry_tab = Frame(self.notebook, bg="#F8FAFC")
        self.events_tab = Frame(self.notebook, bg="#F8FAFC")
        self.history_tab = Frame(self.notebook, bg="#F8FAFC")
        self.scoring_tab = Frame(self.notebook, bg="#F8FAFC")
        tabs = {
            "Today's Check-In": self.entry_tab,
            "Review": self.dashboard,
            "History / Manage Entries": self.history_tab,
            "Treatment Events": self.events_tab,
            "How Scoring Works": self.scoring_tab,
        }
        for tab_name in PRIMARY_TAB_ORDER:
            self.notebook.add(tabs[tab_name], text=tab_name)
        self.build_dashboard()
        self.build_entry_tab()
        self.build_history_tab()
        self.build_events_tab()
        self.build_scoring_tab()
        self.notebook.select(self.entry_tab)
        self.bind_all("<F5>", self.refresh_from_shortcut)

    def build_dashboard(self):
        top = Frame(self.dashboard, bg="#F8FAFC")
        top.pack(fill="x", pady=(0, 10))
        Label(top, text="Review", bg="#F8FAFC", fg="#172033", font=("Segoe UI", 16, "bold")).pack(side=LEFT)
        actions = Frame(top, bg="#F8FAFC")
        actions.pack(side=RIGHT)
        action_commands = {
            "Refresh": self.refresh_all,
            "Import Spreadsheet": self.import_file,
            "Generate PDF": self.create_report,
            "Analysis Workbook": self.export_analysis_file,
        }
        for index, label in enumerate(REVIEW_ACTION_LABELS):
            Button(actions, text=label, command=action_commands[label], width=18).pack(
                side=LEFT,
                padx=(0 if index == 0 else 8, 0),
                pady=2,
            )

        summary = Frame(self.dashboard, bg="#FFFFFF", padx=14, pady=12, highlightthickness=1, highlightbackground="#CBD5E1")
        summary.pack(fill="x", pady=(0, 10))
        Label(summary, text="What stands out", bg="#FFFFFF", fg="#172033", font=("Segoe UI", 12, "bold")).pack(anchor="w")
        self.review_summary = Label(summary, text="No check-ins are available yet.", bg="#FFFFFF", fg="#334155", wraplength=1060, justify=LEFT)
        self.review_summary.pack(anchor="w", pady=(6, 6))
        self.review_highlights = []
        for _ in range(4):
            label = Label(summary, text="", bg="#FFFFFF", fg="#334155", wraplength=1040, justify=LEFT)
            label.pack(anchor="w", pady=2)
            self.review_highlights.append(label)
        Label(
            summary,
            text=CHART_INTERACTION_HELP,
            bg="#FFFFFF",
            fg="#475569",
            wraplength=1040,
            justify=LEFT,
            font=("Segoe UI", 9),
        ).pack(anchor="w", pady=(8, 0))

        self.review_notebook = ttk.Notebook(self.dashboard)
        self.review_notebook.pack(fill=BOTH, expand=True)
        overview = Frame(self.review_notebook, bg="#F8FAFC")
        treatment = Frame(self.review_notebook, bg="#F8FAFC")
        long_term = Frame(self.review_notebook, bg="#F8FAFC")
        self.review_notebook.add(overview, text="Recent Trends")
        self.review_notebook.add(treatment, text="Treatment Cycles")
        self.review_notebook.add(long_term, text="Long-Term Trends")

        self.phq_recent_chart = LineChart(overview, width=520, height=250)
        self.phq_recent_chart.pack(side=LEFT, fill=BOTH, expand=True, padx=(0, 5), pady=6)
        self.gad_recent_chart = LineChart(overview, width=520, height=250)
        self.gad_recent_chart.pack(side=RIGHT, fill=BOTH, expand=True, padx=(5, 0), pady=6)

        self.cycle_labels = []
        self.cycle_charts = []
        for side in (LEFT, RIGHT):
            cycle_frame = Frame(treatment, bg="#F8FAFC")
            cycle_frame.pack(side=side, fill=BOTH, expand=True, padx=6, pady=6)
            label = Label(cycle_frame, text="", bg="#F8FAFC", fg="#334155", wraplength=500, justify=LEFT)
            label.pack(anchor="w", pady=(0, 5))
            chart = LineChart(cycle_frame, width=500, height=220)
            chart.pack(fill=BOTH, expand=True)
            self.cycle_labels.append(label)
            self.cycle_charts.append(chart)

        self.phq_long_chart = LineChart(long_term, width=1050, height=220)
        self.phq_long_chart.pack(fill=BOTH, expand=True, padx=6, pady=(6, 3))
        self.gad_long_chart = LineChart(long_term, width=1050, height=220)
        self.gad_long_chart.pack(fill=BOTH, expand=True, padx=6, pady=(3, 6))

    def build_entry_tab(self):
        form = LabelFrame(self.entry_tab, text="Today's Check-In", bg="#F8FAFC", padx=12, pady=12)
        form.pack(fill=BOTH, expand=True, padx=4, pady=4)
        Button(form, text="← Previous Day", command=lambda: self.navigate_checkin_date(-1)).grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.date_var = StringVar(value=date.today().isoformat())
        self.checkin_date_entry = Entry(form, textvariable=self.date_var, width=16, justify="center")
        self.checkin_date_entry.grid(row=0, column=1, sticky="w", padx=8, pady=4)
        self.checkin_date_entry.bind("<Return>", self.load_checkin_date_from_shortcut)
        self.checkin_date_entry.bind("<FocusOut>", self.auto_load_checkin_date_if_safe)
        self.next_day_button = Button(form, text="Next Day →", command=lambda: self.navigate_checkin_date(1))
        self.next_day_button.grid(row=0, column=2, sticky="w", padx=8)
        Button(form, text="Load Date", command=self.load_checkin_date).grid(row=0, column=3, sticky="w", padx=8)
        self.checkin_mode = Label(form, text="New daily entry", bg="#F8FAFC", fg="#047857", font=("Segoe UI", 10, "bold"))
        self.checkin_mode.grid(row=0, column=4, sticky="w", padx=8)

        self.assessment_item_vars = {}
        assessment_area = Frame(form, bg="#F8FAFC")
        assessment_area.grid(row=1, column=0, columnspan=5, sticky="nsew", pady=(8, 4))
        for col_idx, assessment_id in enumerate(ASSESSMENT_ORDER):
            definition = ASSESSMENTS[assessment_id]
            box = LabelFrame(assessment_area, text=definition.display_name, bg="#F8FAFC", padx=10, pady=10)
            box.grid(row=0, column=col_idx, sticky="nsew", padx=(0, 10))
            self.assessment_item_vars[assessment_id] = []
            for idx, label in enumerate(definition.item_labels, start=1):
                Label(box, text=f"{idx}. {label}", bg="#F8FAFC", wraplength=410, justify=LEFT).grid(row=idx, column=0, sticky="w", pady=3)
                var = IntVar(value=0)
                self.assessment_item_vars[assessment_id].append(var)
                ttk.Spinbox(box, from_=0, to=3, textvariable=var, width=5).grid(row=idx, column=1, sticky="w", padx=8)
        assessment_area.columnconfigure(0, weight=1)
        assessment_area.columnconfigure(1, weight=1)

        Button(form, text="Save Today's Check-In", command=self.save_entry).grid(row=2, column=1, sticky="w", padx=8, pady=(10, 8))
        Label(
            form,
            text="Mindful check-ins intentionally require each symptom to be considered individually; previous responses are never copied or autofilled.",
            bg="#F8FAFC",
            fg="#475569",
            wraplength=900,
            justify=LEFT,
        ).grid(row=3, column=0, columnspan=5, sticky="w", pady=(2, 8))

        self.checkin_details = LabelFrame(
            form,
            text="Anything else to record about today?",
            bg="#F8FAFC",
            padx=10,
            pady=10,
        )
        self.checkin_details.grid(row=4, column=0, columnspan=5, sticky="we", pady=(4, 2))
        Label(
            self.checkin_details,
            text="These details are optional. Your core check-in has already been recorded.",
            bg="#F8FAFC",
            fg="#475569",
        ).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 6))

        Label(self.checkin_details, text="Personal note (optional)", bg="#F8FAFC").grid(row=1, column=0, sticky="nw", pady=(4, 2))
        self.notes_box = Text(self.checkin_details, height=4, width=72)
        self.notes_box.grid(row=1, column=1, columnspan=3, sticky="we", padx=8, pady=(4, 2))
        Label(self.checkin_details, text="Optional note tag", bg="#F8FAFC").grid(row=2, column=0, sticky="w", pady=(4, 2))
        self.note_tag = StringVar(value="")
        ttk.Combobox(
            self.checkin_details,
            textvariable=self.note_tag,
            values=["", "Finances", "Work", "Family", "Health", "Sleep", "Relationships", "Other"],
            state="readonly",
            width=20,
        ).grid(row=2, column=1, sticky="w", padx=8, pady=(4, 2))

        events_box = LabelFrame(self.checkin_details, text="Treatment or health-related events", bg="#F8FAFC", padx=8, pady=6)
        events_box.grid(row=3, column=0, columnspan=4, sticky="we", pady=(10, 2))
        self.checkin_event_vars = {}
        for idx, event_type in enumerate(TREATMENT_EVENT_TYPES):
            var = IntVar(value=0)
            self.checkin_event_vars[event_type] = var
            Checkbutton(events_box, text=event_type, variable=var, bg="#F8FAFC").grid(row=idx // 4, column=idx % 4, sticky="w", padx=6)
        self.include_custom_event = IntVar(value=0)
        Checkbutton(events_box, text="Add Custom Event", variable=self.include_custom_event, bg="#F8FAFC").grid(row=2, column=0, sticky="w", padx=6)
        self.custom_event_type = StringVar(value="")
        Entry(events_box, textvariable=self.custom_event_type, width=28).grid(row=2, column=1, sticky="w", padx=6)
        Label(self.checkin_details, text="Event description (optional)", bg="#F8FAFC").grid(row=4, column=0, sticky="nw", pady=(6, 0))
        self.checkin_event_desc = Text(self.checkin_details, height=2, width=72)
        self.checkin_event_desc.grid(row=4, column=1, columnspan=3, sticky="we", padx=8, pady=(6, 0))
        Button(self.checkin_details, text="Save Optional Details", command=self.save_optional_details).grid(row=5, column=1, sticky="w", padx=8, pady=10)
        Button(self.checkin_details, text="Not Right Now", command=self.hide_checkin_details).grid(row=5, column=2, sticky="w", padx=8, pady=10)
        Button(self.checkin_details, text="Add Another Treatment Event", command=lambda: self.open_history_for_date(self.date_var.get())).grid(row=5, column=3, sticky="e", padx=8, pady=10)
        self.checkin_details.columnconfigure(3, weight=1)
        self.checkin_details.grid_remove()
        form.columnconfigure(4, weight=1)

    def build_history_tab(self):
        top = Frame(self.history_tab, bg="#F8FAFC")
        top.pack(fill="x", padx=6, pady=6)
        self.history_previous_button = Button(top, text="← Previous Day", command=lambda: self.navigate_history_date(-1))
        self.history_previous_button.pack(side=LEFT, padx=(0, 8))
        Label(top, text="Manage records for date (YYYY-MM-DD)", bg="#F8FAFC").pack(side=LEFT)
        self.history_date = StringVar(value=date.today().isoformat())
        self.history_date_entry = Entry(top, textvariable=self.history_date, width=16)
        self.history_date_entry.pack(side=LEFT, padx=8)
        self.history_date_entry.bind("<Return>", self.load_history_date_from_shortcut)
        self.history_date_entry.bind("<FocusOut>", self.auto_load_history_date_if_safe)
        Button(top, text="Load Date", command=self.load_history_date).pack(side=LEFT)
        self.history_next_button = Button(top, text="Next Day →", command=lambda: self.navigate_history_date(1))
        self.history_next_button.pack(side=LEFT, padx=8)
        self.history_status = Label(top, text="", bg="#F8FAFC", fg="#334155")
        self.history_status.pack(side=LEFT, padx=12)

        self.history_assessment_vars = {}
        self.history_assessment_ids = {}
        self.history_assessment_buttons = {}
        self.history_assessment_widgets = {}
        assessments_frame = Frame(self.history_tab, bg="#F8FAFC")
        assessments_frame.pack(fill="x", padx=6, pady=4)
        for col_idx, assessment_id in enumerate(ASSESSMENT_ORDER):
            definition = ASSESSMENTS[assessment_id]
            box = LabelFrame(assessments_frame, text=f"{definition.display_name} responses", bg="#F8FAFC", padx=8, pady=8)
            box.grid(row=0, column=col_idx, sticky="nsew", padx=(0, 8))
            self.history_assessment_vars[assessment_id] = []
            self.history_assessment_widgets[assessment_id] = []
            for idx in range(definition.item_count):
                Label(box, text=f"{idx + 1}", bg="#F8FAFC").grid(row=0, column=idx, padx=2)
                var = IntVar(value=0)
                self.history_assessment_vars[assessment_id].append(var)
                spinner = ttk.Spinbox(box, from_=0, to=3, textvariable=var, width=3)
                spinner.grid(row=1, column=idx, padx=2)
                self.history_assessment_widgets[assessment_id].append(spinner)
            update_button = Button(box, text="Update Existing Entry", command=lambda aid=assessment_id: self.save_history_assessment(aid))
            update_button.grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 0))
            delete_button = Button(box, text="Delete Assessment", command=lambda aid=assessment_id: self.delete_history_assessment(aid))
            delete_button.grid(row=2, column=4, columnspan=5, sticky="e", pady=(8, 0))
            self.history_assessment_buttons[assessment_id] = (update_button, delete_button)
            assessments_frame.columnconfigure(col_idx, weight=1)

        notes_box = LabelFrame(self.history_tab, text="Daily note", bg="#F8FAFC", padx=8, pady=8)
        notes_box.pack(fill="x", padx=6, pady=4)
        self.history_notes = Text(notes_box, height=3, width=80)
        self.history_notes.grid(row=0, column=0, columnspan=3, sticky="we")
        self.history_note_tag = StringVar(value="")
        self.history_note_tag_widget = ttk.Combobox(notes_box, textvariable=self.history_note_tag, values=["", "Finances", "Work", "Family", "Health", "Sleep", "Relationships", "Other"], state="readonly", width=20)
        self.history_note_tag_widget.grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.history_note_update_button = Button(notes_box, text="Update Existing Note", command=self.save_history_note)
        self.history_note_update_button.grid(row=1, column=1, padx=8, pady=(6, 0))
        self.history_note_delete_button = Button(notes_box, text="Delete Note", command=self.delete_history_note)
        self.history_note_delete_button.grid(row=1, column=2, pady=(6, 0))
        notes_box.columnconfigure(0, weight=1)

        event_box = LabelFrame(self.history_tab, text="Treatment events", bg="#F8FAFC", padx=8, pady=8)
        event_box.pack(fill=BOTH, expand=True, padx=6, pady=4)
        self.history_events_table = ttk.Treeview(event_box, columns=["Type", "Description"], show="headings", height=6)
        self.history_events_table.heading("Type", text="Type")
        self.history_events_table.heading("Description", text="Description")
        self.history_events_table.column("Type", width=220)
        self.history_events_table.column("Description", width=650)
        self.history_events_table.grid(row=0, column=0, columnspan=4, sticky="nsew")
        self.history_events_table.bind("<<TreeviewSelect>>", self.select_history_event)
        self.history_event_type = StringVar(value="Therapy")
        self.history_event_type_widget = ttk.Combobox(event_box, textvariable=self.history_event_type, values=[*TREATMENT_EVENT_TYPES, "Custom event"], width=28)
        self.history_event_type_widget.grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.history_event_desc = StringVar(value="")
        self.history_event_desc_widget = Entry(event_box, textvariable=self.history_event_desc, width=65)
        self.history_event_desc_widget.grid(row=1, column=1, sticky="we", padx=6, pady=(6, 0))
        self.history_event_add_button = Button(event_box, text="Add Another Treatment Event", command=self.add_history_event)
        self.history_event_add_button.grid(row=1, column=2, padx=4, pady=(6, 0))
        self.history_event_update_button = Button(event_box, text="Update Selected", command=self.update_history_event)
        self.history_event_update_button.grid(row=2, column=2, padx=4, pady=(6, 0))
        self.history_event_delete_button = Button(event_box, text="Delete Selected", command=self.delete_history_event)
        self.history_event_delete_button.grid(row=2, column=3, padx=4, pady=(6, 0))
        event_box.columnconfigure(1, weight=1)
        event_box.rowconfigure(0, weight=1)

    def build_events_tab(self):
        form = LabelFrame(self.events_tab, text="Medication/Treatment Event Marker", bg="#F8FAFC", padx=12, pady=12)
        form.pack(fill="x", padx=4, pady=4)
        self.event_date = StringVar(value=date.today().isoformat())
        self.event_type = StringVar(value="Ketamine infusion")
        Label(form, text="Date (YYYY-MM-DD)", bg="#F8FAFC").grid(row=0, column=0, sticky="w")
        Entry(form, textvariable=self.event_date, width=16).grid(row=0, column=1, sticky="w", padx=8)
        Label(form, text="Type", bg="#F8FAFC").grid(row=0, column=2, sticky="w", padx=(18, 0))
        ttk.Combobox(
            form,
            textvariable=self.event_type,
            values=[*TREATMENT_EVENT_TYPES, "Custom event"],
            width=30,
        ).grid(row=0, column=3, sticky="w", padx=8)
        Label(form, text="Description", bg="#F8FAFC").grid(row=1, column=0, sticky="nw", pady=(8, 0))
        self.event_desc = Text(form, height=4, width=80)
        self.event_desc.grid(row=1, column=1, columnspan=3, sticky="we", padx=8, pady=(8, 0))
        Button(form, text="Add Another Treatment Event", command=self.save_event).grid(row=2, column=1, sticky="w", padx=8, pady=10)

        self.events_table = ttk.Treeview(self.events_tab, columns=["Date", "Type", "Description"], show="headings", height=16)
        for col, width in [("Date", 120), ("Type", 180), ("Description", 620)]:
            self.events_table.heading(col, text=col)
            self.events_table.column(col, width=width, anchor="w")
        self.events_table.pack(fill=BOTH, expand=True, padx=4, pady=10)

    def build_scoring_tab(self):
        Label(self.scoring_tab, text="How Scoring Works", bg="#F8FAFC", fg="#172033", font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=12, pady=(12, 6))
        explanation = Text(self.scoring_tab, wrap="word", bg="#FFFFFF", fg="#172033", font=("Segoe UI", 10), padx=12, pady=12)
        explanation.insert("1.0", SCORING_EXPLANATION)
        explanation.config(state="disabled")
        explanation.pack(fill=BOTH, expand=True, padx=12, pady=(0, 12))

    def show_scoring_help(self):
        self.notebook.select(self.scoring_tab)

    def refresh_from_shortcut(self, _event=None) -> str:
        """Refresh local Review data using the conventional F5 shortcut."""
        self.refresh_all()
        return "break"

    def show_checkin_details(self):
        self.checkin_details.grid()

    def hide_checkin_details(self):
        self.checkin_details.grid_remove()
        messagebox.showinfo("Today's Check-In", "Today's check-in is recorded.")

    def checkin_state(self) -> tuple:
        return (
            tuple(tuple(var.get() for var in self.assessment_item_vars[assessment_id]) for assessment_id in ASSESSMENT_ORDER),
            self.notes_box.get("1.0", END).strip(),
            self.note_tag.get().strip(),
            tuple((event_type, var.get()) for event_type, var in self.checkin_event_vars.items()),
            self.include_custom_event.get(),
            self.custom_event_type.get().strip(),
            self.checkin_event_desc.get("1.0", END).strip(),
        )

    def has_unsaved_checkin_changes(self) -> bool:
        return self._checkin_snapshot is not None and self.checkin_state() != self._checkin_snapshot

    def confirm_discard_checkin_changes(self) -> bool:
        if not self.has_unsaved_checkin_changes():
            return True
        return messagebox.askyesno(
            "Discard unsaved check-in changes?",
            "You have unsaved check-in changes. Choose Yes to discard them and load another date, "
            "or No to keep editing this date.",
        )

    def navigate_checkin_date(self, days: int):
        if not self.confirm_discard_checkin_changes():
            return
        try:
            target = shift_calendar_date(self.date_var.get(), days)
        except ValueError as exc:
            messagebox.showinfo("Date unavailable", str(exc))
            return
        self.date_var.set(target)
        self.load_checkin_date(confirm_unsaved=False)

    def load_checkin_date_from_shortcut(self, _event=None) -> str:
        self.load_checkin_date()
        return "break"

    def auto_load_checkin_date_if_safe(self, _event=None) -> None:
        """Load a newly entered date on focus change only when no edits could be lost."""
        entry_date = parse_date(self.date_var.get())
        if not entry_date or entry_date == self._loaded_checkin_date:
            return
        if datetime.fromisoformat(entry_date).date() > date.today():
            return
        if self.has_unsaved_checkin_changes():
            self.checkin_mode.config(
                text="Date changed — choose Load Date; current unsaved edits are still preserved.",
                fg="#B45309",
            )
            return
        self.load_checkin_date(confirm_unsaved=False)

    def load_checkin_date(self, confirm_unsaved: bool = True):
        if confirm_unsaved and not self.confirm_discard_checkin_changes():
            return
        entry_date = parse_date(self.date_var.get())
        if not entry_date:
            messagebox.showerror("Invalid date", "Enter the date as YYYY-MM-DD.")
            return
        if datetime.fromisoformat(entry_date).date() > date.today():
            messagebox.showerror("Future date", "Future check-ins are not available.")
            return
        self.date_var.set(entry_date)
        day_data = fetch_day_data(entry_date)
        assessments = day_data["assessments"]
        has_data = bool(assessments or day_data["notes"] or day_data["events"])
        for assessment_id in ASSESSMENT_ORDER:
            row = assessments.get(assessment_id)
            for idx, var in enumerate(self.assessment_item_vars[assessment_id]):
                var.set(row.items[idx] if row else 0)
        self.notes_box.delete("1.0", END)
        self.notes_box.insert("1.0", day_data["notes"])
        self.note_tag.set(day_data["note_tag"])
        existing_types = {event[2] for event in day_data["events"]}
        for event_type, var in self.checkin_event_vars.items():
            var.set(1 if event_type in existing_types else 0)
        custom_events = [event for event in day_data["events"] if event[2] not in TREATMENT_EVENT_TYPES]
        self.include_custom_event.set(1 if custom_events else 0)
        self.custom_event_type.set(custom_events[0][2] if custom_events else "")
        self.checkin_event_desc.delete("1.0", END)
        self.checkin_mode.config(
            text=(
                f"Editing existing records for {entry_date}"
                if has_data
                else f"New check-in for {entry_date}"
            ),
            fg="#B45309" if has_data else "#047857",
        )
        self.next_day_button.config(state="disabled" if entry_date == date.today().isoformat() else "normal")
        self._loaded_checkin_date = entry_date
        self._checkin_snapshot = self.checkin_state()

    def open_history_for_date(self, raw_date: str):
        if not self.confirm_discard_history_changes():
            return
        entry_date = parse_date(raw_date)
        if entry_date:
            self.history_date.set(entry_date)
        self.notebook.select(self.history_tab)
        self.load_history_date(confirm_unsaved=False)

    def history_state(self) -> tuple:
        return (
            tuple(
                (assessment_id, tuple(var.get() for var in self.history_assessment_vars[assessment_id]))
                for assessment_id in ASSESSMENT_ORDER
            ),
            self.history_notes.get("1.0", END).strip(),
            self.history_note_tag.get().strip(),
            self.history_event_type.get().strip(),
            self.history_event_desc.get().strip(),
        )

    def has_unsaved_history_changes(self) -> bool:
        return self._history_snapshot is not None and self.history_state() != self._history_snapshot

    def confirm_discard_history_changes(self) -> bool:
        if not self.has_unsaved_history_changes():
            return True
        return messagebox.askyesno(
            "Discard unsaved History changes?",
            "You have unsaved History changes. Choose Yes to discard them and load another date, "
            "or No to keep editing the currently loaded date.",
        )

    def navigate_history_date(self, days: int):
        if not self.confirm_discard_history_changes():
            return
        try:
            target = shift_calendar_date(self.history_date.get(), days)
        except ValueError as exc:
            messagebox.showinfo("Date unavailable", str(exc))
            return
        self.history_date.set(target)
        self.load_history_date(confirm_unsaved=False)

    def load_history_date_from_shortcut(self, _event=None) -> str:
        self.load_history_date()
        return "break"

    def auto_load_history_date_if_safe(self, _event=None) -> None:
        """Load a newly entered History date only when no unsaved edits could be lost."""
        try:
            entry_date = validate_nonfuture_date(self.history_date.get())
        except ValueError:
            return
        if entry_date == self._loaded_history_date:
            return
        if self.has_unsaved_history_changes():
            self.history_status.config(
                text="Date changed. Choose Load Date to continue; edits for the currently loaded date are preserved."
            )
            return
        self.load_history_date(confirm_unsaved=False)

    def _loaded_history_action_date(self) -> str | None:
        try:
            entry_date = validate_nonfuture_date(self.history_date.get())
        except ValueError as exc:
            messagebox.showerror("Date unavailable", str(exc))
            return None
        if entry_date != self._loaded_history_date:
            messagebox.showinfo(
                "Load the selected date first",
                "The date field differs from the records currently shown. Load the selected date before making changes.",
            )
            return None
        return entry_date

    def load_history_date(self, confirm_unsaved: bool = True):
        if confirm_unsaved and not self.confirm_discard_history_changes():
            return
        try:
            entry_date = validate_nonfuture_date(self.history_date.get())
        except ValueError as exc:
            messagebox.showerror("Date unavailable", str(exc))
            return
        self.history_date.set(entry_date)
        day_data = fetch_day_data(entry_date)
        assessments = day_data["assessments"]
        for assessment_id in ASSESSMENT_ORDER:
            row = assessments.get(assessment_id)
            self.history_assessment_ids[assessment_id] = row.id if row else None
            for idx, var in enumerate(self.history_assessment_vars[assessment_id]):
                var.set(row.items[idx] if row else 0)
        self.history_notes.config(state="normal")
        self.history_notes.delete("1.0", END)
        self.history_notes.insert("1.0", day_data["notes"])
        self.history_note_tag.set(day_data["note_tag"])
        for item in self.history_events_table.get_children():
            self.history_events_table.delete(item)
        for event_id, _event_date, event_type, description in day_data["events"]:
            self.history_events_table.insert("", END, iid=str(event_id), values=[event_type, description])
        self.history_event_type.set("Therapy")
        self.history_event_desc.set("")
        has_records = history_record_count(day_data) > 0
        has_note = bool(day_data["notes"] or day_data["note_tag"])
        for assessment_id, (update_button, delete_button) in self.history_assessment_buttons.items():
            state = "normal" if assessment_id in assessments else "disabled"
            update_button.config(state=state)
            delete_button.config(state=state)
            for spinner in self.history_assessment_widgets[assessment_id]:
                spinner.config(state=state)
        note_edit_state = "normal" if has_records else "disabled"
        self.history_notes.config(state=note_edit_state)
        self.history_note_tag_widget.config(state="readonly" if has_records else "disabled")
        self.history_note_update_button.config(
            state=note_edit_state,
            text="Update Existing Note" if has_note else "Add Daily Note",
        )
        self.history_note_delete_button.config(state="normal" if has_note else "disabled")
        self.history_event_type_widget.config(state="normal" if has_records else "disabled")
        self.history_event_desc_widget.config(state="normal" if has_records else "disabled")
        self.history_event_add_button.config(state="normal" if has_records else "disabled")
        self.history_event_update_button.config(state="normal" if day_data["events"] else "disabled")
        self.history_event_delete_button.config(state="normal" if day_data["events"] else "disabled")
        self.history_status.config(text=history_status_text(entry_date, day_data))
        self.history_next_button.config(state="disabled" if entry_date == date.today().isoformat() else "normal")
        self._loaded_history_date = entry_date
        self._history_snapshot = self.history_state()

    def save_history_assessment(self, assessment_id: str):
        if not self._loaded_history_action_date():
            return
        entry_id = self.history_assessment_ids.get(assessment_id)
        if not entry_id:
            messagebox.showinfo("No existing assessment", "Use Today's Check-In to create a new assessment for this date.")
            return
        items = [int(var.get()) for var in self.history_assessment_vars[assessment_id]]
        try:
            update_assessment_entry(entry_id, assessment_id, items)
            self.refresh_all()
            self.load_history_date(confirm_unsaved=False)
            messagebox.showinfo("Updated", f"Updated the existing {ASSESSMENTS[assessment_id].display_name} assessment without changing its record ID.")
        except Exception as exc:
            messagebox.showerror("Update failed", str(exc))

    def delete_history_assessment(self, assessment_id: str):
        if not self._loaded_history_action_date():
            return
        entry_id = self.history_assessment_ids.get(assessment_id)
        if not entry_id:
            return
        name = ASSESSMENTS[assessment_id].display_name
        if not messagebox.askyesno(
            "Permanently delete assessment?",
            f"Delete the {name} assessment for {self.history_date.get()}? This cannot be undone.",
        ):
            return
        delete_assessment_entry(entry_id, assessment_id)
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def save_history_note(self):
        entry_date = self._loaded_history_action_date()
        if not entry_date:
            return
        update_daily_note(entry_date, self.history_notes.get("1.0", END).strip(), self.history_note_tag.get().strip())
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)
        messagebox.showinfo("Daily note updated", f"Updated the daily note for {entry_date} without creating a duplicate.")

    def delete_history_note(self):
        entry_date = self._loaded_history_action_date()
        if not entry_date or not messagebox.askyesno(
            "Permanently delete daily note?",
            f"Delete the daily note for {entry_date}? This cannot be undone.",
        ):
            return
        delete_daily_note(entry_date)
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def select_history_event(self, _event=None):
        selected = self.history_events_table.selection()
        if not selected:
            return
        already_unsaved = self.has_unsaved_history_changes()
        values = self.history_events_table.item(selected[0], "values")
        self.history_event_type.set(values[0])
        self.history_event_desc.set(values[1])
        if not already_unsaved:
            self._history_snapshot = self.history_state()

    def add_history_event(self):
        entry_date = self._loaded_history_action_date()
        event_type = self.history_event_type.get().strip()
        if not entry_date or not event_type:
            messagebox.showerror("Missing information", "Enter a valid date and event type.")
            return
        add_event(entry_date, event_type, self.history_event_desc.get().strip(), dedupe=False)
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def update_history_event(self):
        selected = self.history_events_table.selection()
        entry_date = self._loaded_history_action_date()
        if not selected or not entry_date:
            messagebox.showinfo("Select an event", "Select the treatment event to update.")
            return
        update_event(int(selected[0]), entry_date, self.history_event_type.get().strip() or "Treatment event", self.history_event_desc.get().strip())
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def delete_history_event(self):
        if not self._loaded_history_action_date():
            return
        selected = self.history_events_table.selection()
        if not selected or not messagebox.askyesno(
            "Permanently delete treatment event?",
            "Delete the selected treatment event? This cannot be undone.",
        ):
            return
        delete_event(int(selected[0]))
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def refresh_all(self):
        self.refresh_events()
        phq_entries = fetch_assessment_entries("phq9")
        gad_entries = fetch_assessment_entries("gad7")
        all_dates = sorted({row.entry_date for row in phq_entries + gad_entries})
        if all_dates:
            latest_date = all_dates[-1]
            self.review_summary.config(text=overall_pattern_summary(phq_entries, gad_entries, latest_date))
            highlights = [
                *symptom_highlights("phq9", phq_entries, latest_date, limit=2),
                *symptom_highlights("gad7", gad_entries, latest_date, limit=2),
            ]
            for index, label in enumerate(self.review_highlights):
                label.config(text=f"• {highlights[index]}" if index < len(highlights) else "")

            phq_recent = phq_entries[-28:]
            gad_recent = gad_entries[-28:]
            self.phq_recent_chart.draw_series(
                phq_recent,
                [("PHQ-9", [row.total for row in phq_recent], "#2563EB")],
                27,
                "Recent PHQ-9 recorded scores",
            )
            self.gad_recent_chart.draw_series(
                gad_recent,
                [("GAD-7", [row.total for row in gad_recent], "#0F766E")],
                21,
                "Recent GAD-7 recorded scores",
            )

            cycles = treatment_cycles(phq_entries, fetch_events(end=latest_date), latest_date)
            for index, (label, chart) in enumerate(zip(self.cycle_labels, self.cycle_charts)):
                if index < len(cycles):
                    cycle = cycles[index]
                    label.config(text=treatment_cycle_observation(cycle))
                    chart.draw_series(
                        cycle.entries,
                        [("PHQ-9", [row.total for row in cycle.entries], "#7C3AED")],
                        27,
                        f"{cycle.label}: recorded PHQ-9 scores",
                    )
                else:
                    label.config(text="A second recorded ketamine infusion is needed to show this cycle.")
                    chart.draw_series([], [], 27, "Treatment cycle not available")

            phq_long = phq_entries[-120:]
            gad_long = gad_entries[-120:]
            self.phq_long_chart.draw_series(
                phq_long,
                [("PHQ-9", [row.total for row in phq_long], "#2563EB")],
                27,
                "Long-term PHQ-9 recorded scores (up to 120 entries)",
            )
            self.gad_long_chart.draw_series(
                gad_long,
                [("GAD-7", [row.total for row in gad_long], "#0F766E")],
                21,
                "Long-term GAD-7 recorded scores (up to 120 entries)",
            )
        else:
            self.review_summary.config(text="No check-ins are available yet. Today's Check-In is ready when you are.")
            for label in self.review_highlights:
                label.config(text="")
            for chart, title, y_max in (
                (self.phq_recent_chart, "Recent PHQ-9 recorded scores", 27),
                (self.gad_recent_chart, "Recent GAD-7 recorded scores", 21),
                (self.phq_long_chart, "Long-term PHQ-9 recorded scores", 27),
                (self.gad_long_chart, "Long-term GAD-7 recorded scores", 21),
            ):
                chart.draw_series([], [], y_max, title)
            for label, chart in zip(self.cycle_labels, self.cycle_charts):
                label.config(text="No recorded ketamine infusion is available for cycle review.")
                chart.draw_series([], [], 27, "Treatment cycle not available")

    def refresh_recent_table(self):
        for row in self.recent_table.get_children():
            self.recent_table.delete(row)
        phq_entries = fetch_assessment_entries("phq9")
        gad_entries = fetch_assessment_entries("gad7")
        phq_by_date = {row.entry_date: row for row in phq_entries}
        gad_by_date = {row.entry_date: row for row in gad_entries}
        dates = sorted(set(phq_by_date) | set(gad_by_date))[-14:]
        prior_phq = None
        prior_gad = None
        for entry_date in dates:
            phq = phq_by_date.get(entry_date)
            gad = gad_by_date.get(entry_date)
            trend_parts = []
            tag = "neutral"
            if phq and prior_phq:
                delta = phq.total - prior_phq.total
                if delta:
                    trend_parts.append(f"PHQ {delta:+d}")
                    tag = "worse" if delta > 0 else "better"
            if gad and prior_gad:
                delta = gad.total - prior_gad.total
                if delta:
                    trend_parts.append(f"GAD {delta:+d}")
                    if tag == "neutral":
                        tag = "worse" if delta > 0 else "better"
            self.recent_table.insert(
                "",
                END,
                values=[
                    entry_date,
                    phq.total if phq else "--",
                    phq.severity if phq else "--",
                    gad.total if gad else "--",
                    gad.severity if gad else "--",
                    ", ".join(trend_parts) or "No change",
                ],
                tags=(tag,),
            )
            prior_phq = phq or prior_phq
            prior_gad = gad or prior_gad

    def refresh_item_table(self):
        for row in self.item_table.get_children():
            self.item_table.delete(row)
        for assessment_id in ASSESSMENT_ORDER:
            definition = ASSESSMENTS[assessment_id]
            entries = fetch_assessment_entries(assessment_id)
            score_14_day = calculate_14_day_symptom_frequency_score(entries, definition.item_count, assessment_id)
            if not entries:
                continue
            for idx, label in enumerate(definition.item_labels, start=1):
                values = [row.items[idx - 1] for row in entries]
                item_score = score_14_day.item_scores[idx - 1]
                days_present = score_14_day.item_counts[idx - 1]
                self.item_table.insert(
                    "",
                    END,
                    values=[definition.display_name, f"Item {idx}: {label}", f"{sum(values) / len(values):.1f}", f"{item_score} ({days_present} days)"],
                )

    def refresh_events(self):
        for row in self.events_table.get_children():
            self.events_table.delete(row)
        for event_id, event_date, event_type, desc in fetch_events():
            self.events_table.insert("", END, iid=str(event_id), values=[event_date, event_type, desc])

    def draw_charts(self):
        return

    def import_file(self):
        path = filedialog.askopenfilename(
            title="Import PHQ-9 spreadsheet",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if not path:
            return
        try:
            if load_workbook is None:
                run_bundled_cli(["--import", path], ("openpyxl",))
                count = "the selected"
            else:
                count = import_spreadsheet(path)
            self.refresh_all()
            messagebox.showinfo("Import complete", f"Imported or updated {count} PHQ-9 entries. Review now shows the refreshed data.")
        except Exception as exc:
            messagebox.showerror("Import failed", str(exc))

    def export_analysis_file(self):
        path = filedialog.asksaveasfilename(
            title="Export analysis-ready workbook",
            initialdir=str(EXPORTS_DIR),
            initialfile=f"Mental_Health_Tracker_Analysis_{date.today().isoformat()}.xlsx",
            defaultextension=".xlsx",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not path:
            return
        try:
            if pd is None:
                run_bundled_cli(["--analysis-export", path], ("pandas", "openpyxl"))
            else:
                export_analysis_workbook(path)
            messagebox.showinfo("Analysis workbook saved", f"Saved the normalized seven-sheet workbook to:\n\n{path}")
        except Exception as exc:
            messagebox.showerror("Analysis export failed", str(exc))

    def save_entry(self):
        entry_date = parse_date(self.date_var.get())
        if not entry_date:
            messagebox.showerror("Invalid date", "Enter the date as YYYY-MM-DD.")
            return
        if datetime.fromisoformat(entry_date).date() > date.today():
            messagebox.showerror("Future date", "Future check-ins are not available.")
            return
        if entry_date != self._loaded_checkin_date:
            messagebox.showinfo(
                "Load the selected date first",
                "The date field differs from the check-in currently shown. Load the selected date before saving "
                "so responses are not applied to the wrong day.",
            )
            return
        existing = fetch_day_data(entry_date)
        saved = []
        for assessment_id in ASSESSMENT_ORDER:
            definition = ASSESSMENTS[assessment_id]
            items = [int(var.get()) for var in self.assessment_item_vars[assessment_id]]
            if any(score < 0 or score > 3 for score in items):
                messagebox.showerror("Invalid score", f"Each {definition.display_name} item must be 0, 1, 2, or 3.")
                return
            if assessment_id == "phq9":
                upsert_entry(
                    entry_date,
                    items,
                    notes=str(existing["notes"]),
                    note_tag=str(existing["note_tag"]),
                    source="manual",
                )
            else:
                upsert_assessment_entry(
                    assessment_id,
                    entry_date,
                    items,
                    notes=str(existing["notes"]),
                    note_tag=str(existing["note_tag"]),
                    source="manual",
                )
            saved.append(definition.display_name)
        self.refresh_all()
        self.load_checkin_date(confirm_unsaved=False)
        self.show_checkin_details()
        action = "Updated existing" if existing["assessments"] else "Saved new"
        messagebox.showinfo(
            "Check-in recorded",
            f"{action} {', '.join(saved)} check-in for {entry_date}. The core check-in is saved. "
            "Optional details can be added now or skipped.",
        )

    def save_optional_details(self):
        entry_date = parse_date(self.date_var.get())
        if not entry_date:
            messagebox.showerror("Invalid date", "Enter the date as YYYY-MM-DD.")
            return
        custom_type = self.custom_event_type.get().strip()
        if self.include_custom_event.get() and not custom_type:
            messagebox.showerror("Missing custom event", "Enter a name for the custom event.")
            return
        notes = self.notes_box.get("1.0", END).strip()
        tag = self.note_tag.get().strip()
        update_daily_note(entry_date, notes, tag)
        description = self.checkin_event_desc.get("1.0", END).strip()
        for event_type, var in self.checkin_event_vars.items():
            if var.get():
                upsert_daily_event(entry_date, event_type, description)
        if self.include_custom_event.get():
            upsert_daily_event(entry_date, custom_type, description)
        self.checkin_event_desc.delete("1.0", END)
        self.refresh_all()
        self.load_checkin_date(confirm_unsaved=False)
        self.hide_checkin_details()

    def save_event(self):
        event_date = parse_date(self.event_date.get())
        if not event_date:
            messagebox.showerror("Invalid date", "Enter the event date as YYYY-MM-DD.")
            return
        add_event(event_date, self.event_type.get().strip() or "Treatment event", self.event_desc.get("1.0", END).strip(), dedupe=False)
        self.event_desc.delete("1.0", END)
        self.refresh_events()
        messagebox.showinfo("Saved", f"Saved event for {event_date}.")

    def create_report(self):
        try:
            start, end = available_report_date_range()
        except ValueError as exc:
            messagebox.showinfo("Report unavailable", str(exc))
            return
        target = filedialog.asksaveasfilename(
            title="Save clinician report PDF",
            initialdir=str(REPORTS_DIR),
            initialfile=f"Mental_Health_Tracker_Report_{start}_to_{end}.pdf",
            defaultextension=".pdf",
            filetypes=[("PDF report", "*.pdf")],
        )
        if not target:
            return
        try:
            if colors is None or PILImage is None:
                run_bundled_cli(
                    ["--report-pdf", target],
                    ("reportlab", "PIL"),
                )
            else:
                generate_report(start, end, target)
            messagebox.showinfo("PDF report saved", f"Saved the full-history clinician discussion report to:\n\n{target}")
        except Exception as exc:
            messagebox.showerror("Report failed", str(exc))


def main():
    parser = argparse.ArgumentParser(description="Local Mental Health Tracker")
    parser.add_argument("--import", dest="import_path", help="Import an Excel workbook, then exit unless --launch is also set.")
    parser.add_argument("--launch", action="store_true", help="Launch the GUI after command-line actions.")
    parser.add_argument("--report-pdf", help="Generate a full-history PDF report at this output path.")
    parser.add_argument("--analysis-export", help="Export a normalized analysis-ready XLSX workbook, then exit.")
    args = parser.parse_args()

    init_db()
    if args.import_path:
        count = import_spreadsheet(args.import_path)
        print(f"Imported or updated {count} entries.")
    if args.report_pdf:
        pdf_path = Path(args.report_pdf)
        start, end = available_report_date_range()
        generate_report(start, end, str(pdf_path))
        print(f"Saved report to {pdf_path}")
    if args.analysis_export:
        export_analysis_workbook(args.analysis_export)
        print(f"Saved analysis workbook to {args.analysis_export}")
    has_cli_action = bool(args.import_path or args.report_pdf or args.analysis_export)
    if args.launch or not has_cli_action:
        PHQ9App().mainloop()


if __name__ == "__main__":
    main()
